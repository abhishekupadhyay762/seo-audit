"""
SEO Audit Toolkit - one page in depth (no screen code in this file)

  audit_page()          everything about one URL: page checks, links, canonical/hreflang targets, details
  page_details()        heading outline, images, JSON-LD objects, text blocks, HTTP headers
  readability()         Flesch reading ease, long sentences and long paragraphs (English pages)
  content_analysis()    most used words and two- and three-word phrases, keyword use
  schema_fields()       required and recommended fields for each schema type Google uses
  mobile_vs_desktop()   the page fetched as a smartphone and as a desktop, with the differences
  score_page()          the on-page score out of 100, the score per area and the checklist
  compare_pages()       your page next to competitor pages, and the phrases they use that you do not
"""

import io
import json
import re
from collections import Counter
from datetime import date
from urllib.parse import urljoin, urlsplit

import pandas as pd
from bs4 import BeautifulSoup

import seo_checks as checks
import seo_core as core
import seo_crawl as crawl
import seo_report as report

# --------------------------------------------------------------------------
# Settings you can safely change
# --------------------------------------------------------------------------
BLOCKER_CAP = 30              # a blocker (noindex, error, robots.txt block, canonical elsewhere) caps the score here
GRADES = [(90, "A"), (75, "B"), (60, "C"), (40, "D"), (0, "F")]
LONG_SENTENCE_WORDS = 25      # sentences longer than this are "long"
LONG_PARAGRAPH_WORDS = 150    # paragraphs longer than this are "long"
FAST_MS, SLOW_MS = 1000, 2500  # server response time bands (measured from this tool's server)
MAX_COMPETITORS = 3

STOP_WORDS = set("""
a about above after again against all also am an and any are aren't as at be because been before being below between
both but by can can't cannot could couldn't did didn't do does doesn't doing don't down during each even every few for
from further get got had hadn't has hasn't have haven't having he he'd he'll he's her here here's hers herself him
himself his how how's however i i'd i'll i'm i've if in into is isn't it it's its itself just let's like make many may
me more most much must mustn't my myself new no nor not now of off on once one only or other ought our ours ourselves
out over own per same shall shan't she she'd she'll she's should shouldn't so some such than that that's the their
theirs them themselves then there there's these they they'd they'll they're they've this those through to too two
under until up upon us use used using very via want was wasn't way we we'd we'll we're we've well were weren't what
what's when when's where where's whether which while who who's whom whose why why's will with within without won't
would wouldn't yet you you'd you'll you're you've your yours yourself yourselves
""".split())

LOCAL_BUSINESS_TYPES = set("""
LocalBusiness ProfessionalService Store Restaurant Dentist MedicalBusiness MedicalClinic Physician LegalService Attorney
Notary HomeAndConstructionBusiness Plumber Electrician HVACBusiness RoofingContractor GeneralContractor HousePainter
Locksmith MovingCompany AutomotiveBusiness AutoRepair AutoDealer FinancialService AccountingService BankOrCreditUnion
InsuranceAgency RealEstateAgent HealthAndBeautyBusiness BeautySalon HairSalon DaySpa NailSalon TattooParlor
FoodEstablishment CafeOrCoffeeShop Bakery BarOrPub FastFoodRestaurant LodgingBusiness Hotel Motel Resort
SportsActivityLocation ExerciseGym TravelAgency EmploymentAgency ChildCare Optician Pharmacy VeterinaryCare
EntertainmentBusiness DryCleaningOrLaundry ShoppingCenter
""".split())
ORGANIZATION_TYPES = {"Organization", "Corporation", "NGO", "EducationalOrganization", "School", "CollegeOrUniversity",
                      "OnlineBusiness", "OnlineStore", "NewsMediaOrganization", "SportsOrganization", "GovernmentOrganization"}

# type -> (required fields, at least one of these, recommended fields); "a|b" means either a or b
SCHEMA_RULES = {
    "Article": ([], [], ["headline", "image", "datePublished", "dateModified", "author.name"]),
    "Product": (["name"], ["offers", "review", "aggregateRating"],
                ["image", "description", "brand", "sku", "offers.price", "offers.priceCurrency", "offers.availability"]),
    "LocalBusiness": (["name", "address"], [], ["telephone", "url", "geo", "openingHoursSpecification|openingHours",
                                                "priceRange"]),
    "Organization": ([], [], ["name", "url", "logo", "sameAs", "contactPoint|telephone"]),
    "WebSite": ([], [], ["name", "url"]),
    "FAQPage": (["mainEntity"], [], []),
    "BreadcrumbList": (["itemListElement"], [], []),
    "Event": (["name", "startDate", "location"], [], ["description", "endDate", "eventStatus", "image", "offers",
                                                       "organizer", "performer"]),
    "Recipe": (["name", "image"], [], ["author", "datePublished", "description", "recipeIngredient",
                                       "recipeInstructions", "totalTime", "recipeYield", "nutrition.calories"]),
    "VideoObject": (["name", "thumbnailUrl", "uploadDate"], [], ["description", "duration", "contentUrl|embedUrl"]),
    "JobPosting": (["title", "description", "datePosted", "hiringOrganization", "jobLocation|applicantLocationRequirements"],
                   [], ["validThrough", "employmentType", "baseSalary"]),
    "Review": (["itemReviewed", "reviewRating.ratingValue", "author"], [], ["datePublished"]),
    "AggregateRating": (["ratingValue", "ratingCount|reviewCount"], [], ["bestRating", "itemReviewed"]),
    "Course": (["name", "description"], [], ["provider"]),
    "SoftwareApplication": (["name", "offers.price"], ["aggregateRating", "review"], ["applicationCategory", "operatingSystem"]),
    "HowTo": ([], [], []),
    "Person": ([], [], ["name", "url", "image", "sameAs", "jobTitle"]),
    "Service": ([], [], ["name", "provider", "areaServed", "serviceType", "description"]),
}
SCHEMA_ALIASES = {"BlogPosting": "Article", "NewsArticle": "Article", "WebApplication": "SoftwareApplication",
                  "MobileApplication": "SoftwareApplication", "ProductGroup": "Product"}
SCHEMA_NOTES = {
    "FAQPage": "Google shows FAQ rich results only for well-known government and health sites; the markup still describes the page.",
    "Article": "Google lists no required fields for Article; the recommended ones help it understand the page.",
    "Course": "Google retired the Course info rich result in 2025; the markup still describes the course.",
    "HowTo": "Google no longer shows HowTo rich results; the markup does no harm.",
    "WebSite": "Google reads the site name from WebSite markup. The sitelinks search box it used to support was retired in 2024.",
    "Person": "No rich result for Person on its own; these fields help Google connect the person to the page.",
    "Service": "No rich result for Service; these fields help describe what is offered.",
    "AggregateRating": "Only for real ratings shown on the page. Google ignores self-serving ratings for LocalBusiness and Organization.",
}


# --------------------------------------------------------------------------
# Page details
# --------------------------------------------------------------------------
def _jsonld_objects(soup):
    objects, errors = [], 0
    for txt in core.jsonld_scripts(soup):
        if not txt:
            continue
        txt = re.sub(r"^\s*(<!--|//\s*<!\[CDATA\[|<!\[CDATA\[)", "", txt)
        txt = re.sub(r"(-->|//\s*\]\]>|\]\]>)\s*$", "", txt).strip()
        try:
            data = json.loads(txt, strict=False)
        except ValueError:
            errors += 1
            continue
        items = data if isinstance(data, list) else [data]
        for item in items:
            if isinstance(item, dict) and isinstance(item.get("@graph"), list):
                objects += [o for o in item["@graph"] if isinstance(o, dict)]
            elif isinstance(item, dict):
                objects.append(item)
    return objects, errors


def page_details(row: dict) -> dict:
    """Heading outline, images, JSON-LD objects, text blocks and response headers of a checked page."""
    x = row["_x"]
    out = {"headings": [], "images": [], "schema_objects": [], "schema_errors": 0, "microdata": [],
           "blocks": [], "headers": x.get("headers", {}), "html_bytes": len(x.get("html", b""))}
    if not x.get("parsed") or not x.get("html"):
        return out
    final = row["Final URL"] or row["URL"]
    soup = BeautifulSoup(x["html"], "lxml")
    base_tag = soup.find("base", href=True)
    try:
        base = urljoin(final, base_tag["href"].strip()) if base_tag else final
    except ValueError:
        base = final

    out["headings"] = [{"Level": h.name.upper(), "Heading": core.node_text(h)}
                       for h in soup.find_all(re.compile(r"^h[1-6]$"))]
    for img in soup.find_all("img"):
        if img.find_parent("noscript") is not None:
            continue
        src = (img.get("src") or img.get("data-src") or img.get("data-lazy-src")
               or (img.get("srcset") or "").split(",")[0].strip().split(" ")[0])
        alt = img.get("alt")
        try:
            src_abs = urljoin(base, src.strip()) if src and not src.startswith("data:") else (src or "")[:60]
        except ValueError:
            src_abs = src
        out["images"].append({
            "Image": src_abs or "(no src)", "Alt": "Missing" if alt is None else ("Empty" if not alt.strip() else "OK"),
            "Alt Text": core.clean(alt) if alt else "", "Width": img.get("width") or "", "Height": img.get("height") or "",
            "Loading": img.get("loading") or "",
        })
    out["schema_objects"], out["schema_errors"] = _jsonld_objects(soup)
    out["microdata"] = core.uniq(core._short_type(t) for el in soup.find_all(attrs={"itemtype": True})
                                 for t in str(el.get("itemtype")).split())

    area = soup.find("main") or soup.find(attrs={"role": "main"}) or soup.find("article") or soup.body or soup
    for el in area.find_all(["p", "li"]):
        if el.find_parent(["nav", "footer", "aside", "header"]) is not None and el.find_parent(["main", "article"]) is None:
            continue
        if el.name == "li" and el.find("p"):
            continue  # the paragraphs inside are counted on their own
        text = core.clean(el.get_text(" "))
        if text:
            out["blocks"].append(text)
    return out


# --------------------------------------------------------------------------
# Readability
# --------------------------------------------------------------------------
_SENTENCE_END = re.compile(r"(?<=[.!?])[\"')\]]*\s+(?=[\"'(\[]?[A-Z0-9])")
_WORD = re.compile(r"[A-Za-z0-9][A-Za-z0-9'’-]*")


def syllables(word: str) -> int:
    w = re.sub(r"[^a-z]", "", word.lower())
    if not w:
        return 0
    if len(w) <= 3:
        return 1
    w = re.sub(r"(?:[^laeiouy]es|ed|[^laeiouy]e)$", "", w)
    w = re.sub(r"^y", "", w)
    return max(1, len(re.findall(r"[aeiouy]{1,2}", w)))


def reading_level(score: float) -> str:
    for limit, label in ((90, "Very easy"), (80, "Easy"), (70, "Fairly easy"), (60, "Standard"),
                         (50, "Fairly difficult"), (30, "Difficult")):
        if score >= limit:
            return label
    return "Very difficult"


def readability(blocks: list, lang: str = "") -> dict:
    """Flesch reading ease (English), sentence and paragraph length, from the page's paragraphs and list items."""
    sentences, words, syl = [], 0, 0
    for block in blocks:
        for sentence in _SENTENCE_END.split(block):
            tokens = _WORD.findall(sentence)
            if tokens:
                sentences.append(tokens)
                words += len(tokens)
                syl += sum(syllables(t) for t in tokens)
    english = (not lang) or lang.lower().startswith("en")
    long_s = [s for s in sentences if len(s) > LONG_SENTENCE_WORDS]
    long_p = [b for b in blocks if len(b.split()) > LONG_PARAGRAPH_WORDS]
    out = {
        "Words": words, "Sentences": len(sentences), "Paragraphs": len(blocks),
        "Average words per sentence": round(words / len(sentences), 1) if sentences else 0,
        "Long sentences": len(long_s),
        "Long sentences (%)": round(100 * len(long_s) / len(sentences), 1) if sentences else 0,
        "Long paragraphs": len(long_p), "Reading ease": None, "Level": "",
        "status": "", "examples": [" ".join(s) for s in long_s[:3]],
    }
    if words < 50:
        out["status"] = "Not enough text to score"
    elif not english:
        out["status"] = f"Not scored: the formula is for English and this page's language is {lang}"
    else:
        score = 206.835 - 1.015 * (words / len(sentences)) - 84.6 * (syl / words)
        out["Reading ease"] = round(max(0.0, min(100.0, score)), 1)
        out["Level"] = reading_level(out["Reading ease"])
    return out


# --------------------------------------------------------------------------
# Content analysis
# --------------------------------------------------------------------------
_CONTENT_WORD = re.compile(r"[^\W\d_][\w'’-]*", re.UNICODE)


def content_words(text: str) -> list:
    return [w.strip("'’-") for w in _CONTENT_WORD.findall((text or "").lower()) if w.strip("'’-")]


_BREAK = re.compile(r"[.!?;:,()\[\]{}|•·–—\n\r\t\u00b6]+")  # \u00b6 = core.BLOCK_MARK, the end of a block


def phrase_words(text: str) -> list:
    """The words of the text with a "|" wherever a sentence or clause ends, so phrases never run across it."""
    out = []
    for part in _BREAK.split(text or ""):
        words = content_words(part)
        if words:
            out.extend(words)
            out.append("|")
    return out


def top_phrases(words: list, n: int, limit: int = 15, min_count: int = 2) -> list:
    grams = Counter(
        " ".join(g) for g in zip(*[words[i:] for i in range(n)])
        if "|" not in g and g[0] not in STOP_WORDS and g[-1] not in STOP_WORDS and len(set(g)) == n
    )
    return [(p, c) for p, c in grams.most_common(limit * 3) if c >= min_count][:limit]


def content_analysis(text: str, html_bytes: int, keyword: str = "", marked_text: str = "") -> dict:
    words = content_words(text)
    total = len(words)
    useful = [w for w in words if w not in STOP_WORDS and len(w) > 2]
    top = Counter(useful).most_common(20)
    marked = phrase_words(marked_text or text)
    phrases = [(p, 2, c) for p, c in top_phrases(marked, 2)] + [(p, 3, c) for p, c in top_phrases(marked, 3)]
    numbers = {
        "Words (main content)": total, "Different words": len(set(words)),
        "Text to HTML (%)": round(100 * len(text) / html_bytes, 1) if html_bytes else 0,
    }
    if keyword:
        count = checks.keyword_count(keyword, text)
        numbers["Keyword uses"] = count
        numbers["Keyword density (%)"] = round(100 * count * len(checks.keyword_tokens(keyword)) / total, 2) if total else 0
    return {
        "numbers": numbers,
        "words": pd.DataFrame([{"Word": w, "Uses": c, "Density (%)": round(100 * c / total, 2) if total else 0}
                               for w, c in top], columns=["Word", "Uses", "Density (%)"]),
        "phrases": pd.DataFrame([{"Phrase": p, "Words": n, "Uses": c} for p, n, c in phrases],
                                columns=["Phrase", "Words", "Uses"]),
        "word_list": words,
        "phrase_list": marked,
    }


# --------------------------------------------------------------------------
# Schema field check
# --------------------------------------------------------------------------
def _present(obj, keys) -> bool:
    if not keys:
        return obj not in (None, "", [], {})
    if isinstance(obj, list):
        return any(_present(o, keys) for o in obj)
    if isinstance(obj, dict) and keys[0] in obj:
        return _present(obj[keys[0]], keys[1:])
    return False


def _has(obj, field: str) -> bool:
    return any(_present(obj, alt.split(".")) for alt in field.split("|"))


def _as_list(value):
    return value if isinstance(value, list) else ([value] if value is not None else [])


def _rule_for(types):
    for t in types:
        t = core._short_type(t)
        key = SCHEMA_ALIASES.get(t, t)
        if key in SCHEMA_RULES:
            return t, key
        if t in LOCAL_BUSINESS_TYPES:
            return t, "LocalBusiness"
        if t in ORGANIZATION_TYPES:
            return t, "Organization"
    return (core._short_type(types[0]) if types else "Unknown"), None


def schema_fields(objects: list, microdata: list = None, errors: int = 0) -> pd.DataFrame:
    """One row per schema item: which required and recommended fields are missing."""
    rows = []
    for obj in objects:
        types = [t for t in _as_list(obj.get("@type")) if isinstance(t, str)]
        if not types:
            continue
        name, key = _rule_for(types)
        if key is None:
            rows.append({"Type": name, "Result": "Not checked", "Missing Required": "", "Missing Recommended": "",
                         "Note": "No Google rich result rules for this type."})
            continue
        required, one_of, recommended = SCHEMA_RULES[key]
        missing_req = [f.replace("|", " or ") for f in required if not _has(obj, f)]
        if one_of and not any(_has(obj, f) for f in one_of):
            missing_req.append("one of " + ", ".join(one_of))
        if key == "FAQPage":
            bad = sum(1 for q in _as_list(obj.get("mainEntity"))
                      if not (isinstance(q, dict) and _has(q, "name") and _has(q, "acceptedAnswer.text")))
            if bad:
                missing_req.append(f"name or acceptedAnswer.text in {bad} question{'s' if bad != 1 else ''}")
        if key == "BreadcrumbList":
            items = _as_list(obj.get("itemListElement"))
            bad = sum(1 for i, it in enumerate(items)
                      if not (isinstance(it, dict) and _has(it, "position") and (_has(it, "name") or _has(it, "item.name"))
                              and (i == len(items) - 1 or _has(it, "item"))))
            if bad:
                missing_req.append(f"position, name or item in {bad} list item{'s' if bad != 1 else ''}")
        missing_rec = [f.replace("|", " or ") for f in recommended if not _has(obj, f)]
        result = "Missing required fields" if missing_req else ("Missing recommended fields" if missing_rec else "OK")
        rows.append({"Type": name, "Result": result, "Missing Required": ", ".join(missing_req),
                     "Missing Recommended": ", ".join(missing_rec), "Note": SCHEMA_NOTES.get(key, "")})
    for t in microdata or []:
        rows.append({"Type": t, "Result": "Not checked", "Missing Required": "", "Missing Recommended": "",
                     "Note": "Microdata: fields are not checked. Test it with Google's Rich Results Test."})
    if errors:
        rows.append({"Type": "(JSON-LD block)", "Result": "Invalid JSON", "Missing Required": "", "Missing Recommended": "",
                     "Note": f"{errors} JSON-LD block{'s' if errors != 1 else ''} could not be read and {'are' if errors != 1 else 'is'} ignored by Google."})
    return pd.DataFrame(rows, columns=["Type", "Result", "Missing Required", "Missing Recommended", "Note"])


# --------------------------------------------------------------------------
# Mobile versus desktop
# --------------------------------------------------------------------------
COMPARE_EXACT = ["Status Code", "Final URL", "Indexable", "Title", "Meta Description", "H1", "Canonical URL",
                 "Meta Robots", "Schema Types", "Hreflang Count"]
COMPARE_SIZE = ["Word Count", "Internal Links", "Images"]
KEY_FIELDS = {"Status Code", "Final URL", "Indexable", "Title", "Meta Description", "H1", "Canonical URL", "Meta Robots"}


def device_sessions(as_googlebot: bool, pool: int = 2) -> dict:
    """One session per device, sending a smartphone or a desktop User-Agent."""
    agents = ({"Smartphone": core.USER_AGENTS["Googlebot Smartphone"], "Desktop": core.USER_AGENTS["Googlebot Desktop"]}
              if as_googlebot else
              {"Smartphone": core.USER_AGENTS["Chrome (mobile)"], "Desktop": core.USER_AGENTS["Chrome (desktop)"]})
    return {name: core.make_session(pool, agent) for name, agent in agents.items()}


def mobile_vs_desktop(url: str, settings: dict, robots, as_googlebot: bool = True, sessions: dict = None):
    """Fetches the page as a smartphone and as a desktop. Returns (table, note).

    sessions  optional {"Smartphone": session, "Desktop": session} to reuse across many pages"""
    sessions = sessions or device_sessions(as_googlebot)
    rows = {}
    for name in ("Smartphone", "Desktop"):
        rows[name] = core.audit_url(sessions[name], url, settings["timeout"], 0, robots, None)
    m, d = rows["Smartphone"], rows["Desktop"]
    table = []

    def text(v):
        return "" if v is None or v is pd.NA else str(v)
    for field in COMPARE_EXACT:
        same = text(m.get(field)) == text(d.get(field))
        table.append({"Field": field, "Smartphone": text(m.get(field)), "Desktop": text(d.get(field)),
                      "Same": "Yes" if same else "No"})
    for field in COMPARE_SIZE:
        a, b = m.get(field) or 0, d.get(field) or 0
        same = abs(a - b) <= max(2, 0.1 * max(a, b))
        table.append({"Field": field, "Smartphone": text(m.get(field)), "Desktop": text(d.get(field)),
                      "Same": "Yes" if same else "No"})
    has_vp = lambda r: "Yes" if r.get("Viewport") else "No"
    table.append({"Field": "Viewport tag", "Smartphone": has_vp(m), "Desktop": has_vp(d),
                  "Same": "Yes" if has_vp(m) == has_vp(d) else "No"})
    note = ""
    statuses = (m["Final Status Code"], d["Final Status Code"])
    if as_googlebot and (None in statuses or any(s in (401, 403, 429) for s in statuses if s)):
        note = ("The site may block requests that claim to be Googlebot. Run the comparison again as Chrome to "
                "see what visitors get.")
    return pd.DataFrame(table, columns=["Field", "Smartphone", "Desktop", "Same"]), note


# --------------------------------------------------------------------------
# The on-page score
# --------------------------------------------------------------------------
def grade(score: int) -> str:
    return next(g for limit, g in GRADES if score >= limit)


def score_page(row: dict, details: dict, links, read: dict, schema: pd.DataFrame, kw: dict, mobile) -> dict:
    """Points per check, grouped by area. Returns score, grade, areas table, checklist and blockers."""
    checks_ = []
    blockers = []

    def add(area, name, earned, possible, found, fix="", blocker=False):
        earned = round(min(max(earned, 0), possible), 2)
        if possible == 0:
            status = "Info"
        elif earned >= possible:
            status = "Pass"
        elif earned == 0:
            status = "Fail"
        else:
            status = "Warning"
        if blocker:
            blockers.append(blocker if isinstance(blocker, str) else name)
        checks_.append({"Area": area, "Check": name, "Result": status, "Found": found, "Points": earned,
                        "Out Of": possible, "How To Fix": fix if status in ("Fail", "Warning") else ""})

    fix = lambda issue: report.advice_for(issue)[1]
    x = row["_x"]
    fs, hops, parsed = row["Final Status Code"], row["Redirect Hops"], x.get("parsed")

    # ---- Indexability (20)
    a = "Indexability"
    if fs is None:
        add(a, "Page answers 200", 0, 6, f"No answer ({row['Fetch Error']})", fix("Fetch failed"),
            f"The page did not answer ({row['Fetch Error']})")
    elif fs >= 400:
        add(a, "Page answers 200", 0, 6, f"Status {fs}", fix("4xx client error" if fs < 500 else "5xx server error"),
            f"The page answers {fs}, not 200")
    elif hops:
        add(a, "Page answers 200", 3, 6, f"Redirects ({row['Status Code']}) to {row['Final URL']}", fix("URL redirects"))
    else:
        add(a, "Page answers 200", 6, 6, str(fs))
    rb = row["Robots.txt"] or "No robots.txt"
    if rb == "Blocked":
        add(a, "Allowed by robots.txt", 0, 4, "Blocked for Googlebot", fix("Blocked by robots.txt"),
            "robots.txt blocks Googlebot from this page")
    elif rb.startswith("Unreachable"):
        add(a, "Allowed by robots.txt", 2, 4, rb, fix("robots.txt could not be read"))
    else:
        add(a, "Allowed by robots.txt", 4, 4, rb)
    noindex = x.get("noindex_meta") or x.get("noindex_header")
    add(a, "No noindex", 0 if noindex else 4, 4, "Noindex found" if noindex else "Indexable", fix("Noindex"),
        ("A noindex " + ("meta tag" if x.get("noindex_meta") else "X-Robots-Tag header") + " tells Google not to index the page")
        if noindex else False)
    cs = row["Canonical Status"]
    if cs == "Self":
        add(a, "Canonical tag", 3, 3, "Points to this page")
    elif cs == "Different":
        add(a, "Canonical tag", 0, 3, f"Points to {row['Canonical URL']}", fix("Canonicalised to another URL"),
            "The canonical tag points to another URL, so Google indexes that one instead")
    elif cs.startswith("Multiple"):
        add(a, "Canonical tag", 1, 3, "Several canonical tags that disagree", fix("Multiple canonical tags"))
    else:
        add(a, "Canonical tag", 1.5 if parsed else 0, 3, "Missing", fix("Canonical missing"))
    https = (row["Final URL"] or row["URL"]).lower().startswith("https://")
    add(a, "HTTPS", 3 if https else 0, 3, "Yes" if https else "No, the page loads over http", fix("Not HTTPS"))

    loaded = fs is not None and fs < 400
    if not loaded:
        add("Other checks", "Content, links, schema and the rest", 0, 0, "Skipped: the page did not load")
    else:
        # ---- Title and meta (15)
        a = "Title & meta"
        title, tlen, tpx = row["Title"], row["Title Length"] or 0, row["Title Pixels"] or 0
        add(a, "Title present", 5 if title else 0, 5, title or "Missing", fix("Title missing"))
        if not title:
            add(a, "Title length", 0, 3, "No title", fix("Title missing"))
        elif tpx > core.TITLE_MAX_PX:
            add(a, "Title length", 1.5, 3, f"{tlen} characters, about {tpx}px (cut after about {core.TITLE_MAX_PX}px)", fix("Title too long"))
        elif tlen < core.TITLE_MIN_CHARS:
            add(a, "Title length", 1.5, 3, f"{tlen} characters", fix("Title too short"))
        else:
            add(a, "Title length", 3, 3, f"{tlen} characters, about {tpx}px")
        desc, dlen = row["Meta Description"], row["Meta Length"] or 0
        add(a, "Meta description present", 4 if desc else 0, 4, desc or "Missing", fix("Meta description missing"))
        if not desc:
            add(a, "Meta description length", 0, 2, "No description", fix("Meta description missing"))
        elif dlen > core.META_MAX_CHARS:
            add(a, "Meta description length", 1, 2, f"{dlen} characters", fix("Meta description too long"))
        elif dlen < core.META_MIN_CHARS:
            add(a, "Meta description length", 1, 2, f"{dlen} characters", fix("Meta description too short"))
        else:
            add(a, "Meta description length", 2, 2, f"{dlen} characters")
        multi = x.get("multi_title") or x.get("multi_meta")
        add(a, "One title and one description tag", 0 if multi else 1, 1, "Duplicates found" if multi else "Yes",
            fix("Multiple title tags"))

        # ---- Headings and content (15)
        a = "Headings & content"
        h1n = row["H1 Count"] or 0
        if h1n == 1 and row["H1"]:
            add(a, "One H1", 4, 4, row["H1"])
        elif h1n > 1:
            add(a, "One H1", 2, 4, f"{h1n} H1 headings", fix("Multiple H1"))
        elif h1n == 1:
            add(a, "One H1", 1, 4, "H1 is empty", fix("H1 empty"))
        else:
            add(a, "One H1", 0, 4, "Missing", fix("H1 missing"))
        h2n = row["H2 Count"] or 0
        add(a, "H2 subheadings", 2 if h2n else 0, 2, f"{h2n} H2 headings",
            "Break the content into sections with descriptive H2 subheadings.")
        order = row["Heading Structure"] or ""
        add(a, "Heading order", 1 if order == "OK" else 0, 1, order or "No headings", fix("Heading levels skipped"))
        words = row["Word Count"] or 0
        add(a, "Content length", 4 if words >= core.THIN_CONTENT_WORDS else (2 if words >= 150 else 0), 4,
            f"{words} words of main content", fix("Thin content"))
        ease = read.get("Reading ease")
        if ease is None:
            add(a, "Readability", 0, 0, read.get("status") or "Not scored")
        else:
            add(a, "Readability", 2 if ease >= 50 else (1 if ease >= 30 else 0), 2, f"Reading ease {ease} ({read['Level']})",
                "Use shorter sentences and everyday words; split long paragraphs.")
        js = row["JS Rendered"] == "Likely"
        add(a, "Content in the HTML", 0 if js else 2, 2, "Looks built by JavaScript" if js else "Yes", fix("Possibly JS-rendered"))

        # ---- Images (10)
        a = "Images"
        imgs = details["images"]
        if not imgs:
            add(a, "Alt attributes", 0, 0, "No images on the page")
        else:
            n = len(imgs)
            with_alt = sum(1 for i in imgs if i["Alt"] != "Missing")
            with_dims = sum(1 for i in imgs if i["Width"] and i["Height"])
            add(a, "Alt attributes", 7 * with_alt / n, 7, f"{n - with_alt} of {n} images have no alt attribute",
                fix("Images missing alt attribute"))
            add(a, "Width and height set", 3 * with_dims / n, 3, f"{n - with_dims} of {n} images have no width or height",
                fix("Images without width/height"))

        # ---- Links (10)
        a = "Links"
        internal = row["Internal Links"] or 0
        add(a, "Internal links", 3 if internal >= 3 else (1.5 if internal else 0), 3, f"{internal} internal links",
            "Link to related pages with descriptive anchor text so visitors and Google can move through the site.")
        if links is None:
            add(a, "No broken links", 0, 0, "Not checked")
            add(a, "No internal links to redirects", 0, 0, "Not checked")
        else:
            broken_int = int(((links["Link Type"] == "Internal") & links["Result"].map(crawl.is_broken)).sum())
            broken_ext = int(((links["Link Type"] == "External") & links["Result"].map(crawl.is_broken)).sum())
            to_redirect = int(((links["Link Type"] == "Internal") & links["Result"].str.startswith("Redirect (")).sum())
            earned = 4 if not (broken_int or broken_ext) else (2 if not broken_int else 0)
            add(a, "No broken links", earned, 4, f"{broken_int} broken internal, {broken_ext} broken external",
                fix("Broken internal links"))
            add(a, "No internal links to redirects", 0 if to_redirect else 1, 1, f"{to_redirect} link{'' if to_redirect == 1 else 's'} to redirects",
                fix("Internal links to redirects"))
        empty = row["Empty Anchor Links"] or 0
        add(a, "Links have anchor text", 2 if not empty else 1, 2, f"{empty} link{'' if empty == 1 else 's'} without text",
            fix("Links with empty anchor text"))

        # ---- Schema and social (10)
        a = "Schema & social"
        add(a, "Schema markup", 3 if row["Schema"] == "Yes" else 0, 3, row["Schema Types"] or "None", fix("No schema markup"))
        if details["schema_errors"]:
            add(a, "Valid JSON-LD", 0, 2, f"{details['schema_errors']} block(s) cannot be read", fix("Invalid JSON-LD schema"))
        elif details["schema_objects"]:
            add(a, "Valid JSON-LD", 2, 2, "Yes")
        else:
            add(a, "Valid JSON-LD", 0, 0, "No JSON-LD on the page")
        checkable = schema[schema["Result"].isin(["OK", "Missing required fields", "Missing recommended fields"])] if len(schema) else schema
        if len(checkable):
            missing = checkable[checkable["Result"] == "Missing required fields"]
            add(a, "Required schema fields", 0 if len(missing) else 2, 2,
                ", ".join(f"{t}: {m}" for t, m in zip(missing["Type"], missing["Missing Required"])) or "All present",
                "Add the missing required fields so the markup is eligible for rich results.")
        else:
            add(a, "Required schema fields", 0, 0, "No schema types with Google rules")
        og = bool(row["OG Title"]) + bool(row["OG Image"])
        add(a, "Open Graph tags", og, 2, f"og:title {'yes' if row['OG Title'] else 'no'}, og:image {'yes' if row['OG Image'] else 'no'}",
            fix("Open Graph tags incomplete"))
        add(a, "Twitter card", 1 if row["Twitter Card"] else 0, 1, row["Twitter Card"] or "Missing",
            "Add a twitter:card meta tag (summary_large_image works for most pages).")

        # ---- Mobile and technical (10)
        a = "Mobile & technical"
        add(a, "Viewport tag", 3 if row["Viewport"] else 0, 3, row["Viewport"] or "Missing", fix("Viewport tag missing"))
        add(a, "Lang attribute", 1 if row["Lang"] else 0, 1, row["Lang"] or "Missing", fix("Lang attribute missing"))
        ms = row["Response Time (ms)"]
        if ms is None:
            add(a, "Server response time", 0, 0, "Not measured")
        else:
            add(a, "Server response time", 2 if ms < FAST_MS else (1 if ms < SLOW_MS else 0), 2,
                f"{ms} ms (measured from this tool's server)",
                "Speed up the server: page caching, a faster host or a CDN.")
        kb = (details["html_bytes"] or 0) / 1024
        if parsed:
            add(a, "HTML size", 1 if kb < 500 else (0.5 if kb < 1500 else 0), 1, f"{kb:.0f} KB",
                "Remove inline code and data from the HTML and load less above the fold.")
        if mobile is None:
            add(a, "Mobile and desktop match", 0, 0, "Not compared")
        else:
            diff = mobile[mobile["Same"] == "No"]["Field"].tolist()
            key_diff = [f for f in diff if f in KEY_FIELDS]
            earned = 3 if not diff else (0 if key_diff else 1.5)
            add(a, "Mobile and desktop match", earned, 3, "Same" if not diff else "Different: " + ", ".join(diff),
                "Google indexes the mobile version. Serve the same title, description, headings, content, links and "
                "tags to phones as to desktops.")

        # ---- Keyword (10, only when a keyword is given)
        if kw:
            a = "Keyword"
            part = {"Yes": 1, "Partial": 0.5, "No": 0, "": 0}
            for col, name, pts, issue in (("In Title", "Keyword in title", 3, "Keyword not in title"),
                                          ("In H1", "Keyword in H1", 2, "Keyword not in H1"),
                                          ("In Meta Description", "Keyword in meta description", 1.5, "Keyword not in meta description"),
                                          ("In URL", "Keyword in URL", 1, "Keyword not in URL"),
                                          ("In First 100 Words", "Keyword in first 100 words", 1.5, "Keyword not in the first 100 words")):
                add(a, name, pts * part[kw.get(col, "")], pts, kw.get(col, "") or "No", fix(issue))
            density, count = kw.get("Keyword Density (%)") or 0, kw.get("Keyword Count") or 0
            add(a, "Keyword use", 1 if (count and density <= checks.KEYWORD_DENSITY_HIGH) else 0, 1,
                f"{count} uses, {density}% density",
                fix("High keyword density") if count else "Use the keyword naturally in the main content.")

    table = pd.DataFrame(checks_, columns=["Area", "Check", "Result", "Found", "Points", "Out Of", "How To Fix"])
    scored = table[table["Out Of"] > 0]
    total = scored["Out Of"].sum()
    score = int(round(100 * scored["Points"].sum() / total)) if total else 0
    if blockers:
        score = min(score, BLOCKER_CAP)
    if not loaded:
        score = 0
    areas = (scored.groupby("Area", sort=False)[["Points", "Out Of"]].sum().reset_index())
    areas["Score (%)"] = (100 * areas["Points"] / areas["Out Of"]).round(0).astype(int)
    areas["Points"] = areas["Points"].round(1)
    return {"score": score, "grade": grade(score), "areas": areas, "checklist": table, "blockers": blockers,
            "counts": table["Result"].value_counts().to_dict(), "loaded": loaded}


# --------------------------------------------------------------------------
# The whole audit of one page
# --------------------------------------------------------------------------
def audit_page(session, url: str, settings: dict, keyword: str = "", check_links: bool = True,
               check_targets: bool = True, mobile: str = "googlebot", on_step=None) -> dict:
    """mobile: "googlebot", "chrome" or "" (skip the mobile versus desktop comparison)."""
    step = on_step or (lambda text: None)
    robots = core.RobotsCache()
    step("Reading the page")
    row = core.audit_url(session, url, settings["timeout"], 0, robots, {"links": True, "text": True, "keep_html": True})
    x = row["_x"]
    details = page_details(row)
    source = row["Final URL"] or row["URL"]
    site = core.host_key(source)
    links_list = [(source, u, anchor, nf, core.host_key(u) == site) for u, anchor, nf in x.pop("links", [])]

    links_df = None
    if check_links and x.get("parsed"):
        step(f"Checking {len(links_list):,} links")
        links_df = crawl.check_links(session, links_list, [row], settings, True)
    if check_targets and x.get("parsed"):
        step("Checking canonical and hreflang targets")
        checks.validate_targets(session, [row], robots, settings)

    kw = None
    if keyword:
        row["Target Keyword"] = keyword
        kw, _ = checks.keyword_extra([row])(row)
    read = readability(details["blocks"], row["Lang"])
    content = content_analysis(x.get("main_text", ""), details["html_bytes"], keyword, x.pop("main_text_marked", ""))
    schema = schema_fields(details["schema_objects"], details["microdata"], details["schema_errors"])
    mobile_df, mobile_note = None, ""
    if mobile and row["Final Status Code"] is not None:
        step("Comparing the smartphone and desktop versions")
        mobile_df, mobile_note = mobile_vs_desktop(url, settings, robots, mobile == "googlebot")
    score = score_page(row, details, links_df, read, schema, kw, mobile_df)
    df = core.finalize([row], core.KEYWORD_GROUPS if keyword else core.GROUPS,
                       checks.keyword_extra([row]) if keyword else None)
    x.pop("html", None)  # no longer needed; keeps memory low
    return {"url": url, "row": row, "df": df, "details": details, "links": links_df, "readability": read,
            "content": content, "schema": schema, "mobile": mobile_df, "mobile_note": mobile_note,
            "keyword": keyword, "keyword_values": kw, "score": score}


# --------------------------------------------------------------------------
# The score (and the smartphone vs desktop result) for many pages at once
# --------------------------------------------------------------------------
def weakest_area(areas: pd.DataFrame) -> str:
    if areas.empty:
        return ""
    low = areas.sort_values(["Score (%)", "Out Of"], ascending=[True, False]).iloc[0]
    return f"{low['Area']} ({int(low['Score (%)'])}%)" if low["Score (%)"] < 100 else ""


def row_hook(score: bool = True, mobile: str = "", settings: dict = None, robots=None) -> dict:
    """Page-check options that add the on-page score and/or the smartphone vs desktop comparison to every row.

    Merge the result into the opts given to core.audit_url / core.run_audit. The work runs in the
    worker thread straight after each page is read, and the page HTML is dropped afterwards.
    mobile: "", "chrome" or "googlebot"
    """
    if not score and not mobile:
        return {}
    sessions = device_sessions(mobile == "googlebot", (settings or {}).get("workers", 4)) if mobile else None

    def after(row: dict):
        x = row["_x"]
        try:
            mob = None
            if mobile and row["Final Status Code"] is not None:
                mob, note = mobile_vs_desktop(row["URL"], settings, robots, mobile == "googlebot", sessions)
                diff = mob[mob["Same"] == "No"]["Field"].tolist()
                row["Mobile Match"] = "Different" if diff else "Same"
                row["Mobile Differences"] = ", ".join(diff) + (f" ({note})" if note else "")
                if diff:
                    key = [f for f in diff if f in KEY_FIELDS]
                    x.setdefault("more_issues", []).append(
                        (core.WARNING if key else core.NOTICE, "Smartphone and desktop versions differ",
                         "Smartphone and desktop differ: " + ", ".join(diff)))
            if score:
                details = page_details(row)
                read = readability(details["blocks"], row["Lang"])
                schema = schema_fields(details["schema_objects"], details["microdata"], details["schema_errors"])
                s = score_page(row, details, None, read, schema, None, mob)
                row["On-page Score"], row["Grade"] = s["score"], s["grade"]
                row["Weakest Area"] = weakest_area(s["areas"]) if s["loaded"] else "Page did not load"
        finally:
            for key in ("html", "headers", "main_text_marked"):
                x.pop(key, None)

    return {"keep_html": bool(score), "after": after}


# --------------------------------------------------------------------------
# Competitor comparison
# --------------------------------------------------------------------------
def _metrics(audit: dict) -> dict:
    r, s = audit["row"], audit["score"]
    areas = {f"{a} score": f"{p:g} / {o:g}" for a, p, o in zip(s["areas"]["Area"], s["areas"]["Points"], s["areas"]["Out Of"])}
    out = {"On-page score": f"{s['score']} ({s['grade']})", **areas,
           "Status": str(r["Status Code"] if r["Status Code"] is not None else r["Fetch Error"]),
           "Title": r["Title"], "Title length": str(r["Title Length"] or 0),
           "Meta description length": str(r["Meta Length"] or 0), "H1": r["H1"],
           "H2 headings": str(r["H2 Count"] or 0), "Words (main content)": str(r["Word Count"] or 0),
           "Reading ease": str(audit["readability"].get("Reading ease") or "-"),
           "Images": str(r["Images"] or 0), "Images missing alt": str(r["Alt Missing"] or 0),
           "Internal links": str(r["Internal Links"] or 0), "External links": str(r["External Links"] or 0),
           "Schema types": r["Schema Types"] or "None"}
    kw = audit.get("keyword_values")
    if kw:
        out.update({"Keyword places (of 5)": kw.get("Keyword Score", ""), "Keyword in title": kw.get("In Title", ""),
                    "Keyword in H1": kw.get("In H1", ""), "Keyword uses": str(kw.get("Keyword Count") or 0)})
    return out


def compare_pages(main: dict, competitors: list) -> dict:
    """Side-by-side table, phrases competitors use that your page does not, and their schema and H2s."""
    pages = [("Your page", main)] + [(f"Competitor {i}", c) for i, c in enumerate(competitors, 1)]
    metrics = [(label, _metrics(a)) for label, a in pages]
    names = list(dict.fromkeys(k for _, m in metrics for k in m))
    table = pd.DataFrame([{"Measure": n, **{label: m.get(n, "Not scored" if n.endswith(" score") else "")
                                             for label, m in metrics}} for n in names])
    table = pd.concat([pd.DataFrame([{"Measure": "URL", **{label: a["url"] for label, a in pages}}]), table],
                      ignore_index=True)

    yours = " " + " ".join(main["content"]["word_list"]) + " "
    used = {}
    for _, comp in pages[1:]:
        words = comp["content"].get("phrase_list") or comp["content"]["word_list"]
        seen = set()
        for n in (2, 3):
            for phrase, count in top_phrases(words, n, limit=40):
                if phrase in seen:
                    continue
                seen.add(phrase)
                entry = used.setdefault(phrase, [0, 0])
                entry[0] += 1
                entry[1] += count
    need = 2 if len(competitors) >= 2 else 1
    gaps = sorted(((p, n, c) for p, (n, c) in used.items() if n >= need and f" {p} " not in yours),
                  key=lambda t: (-t[1], -t[2]))[:30]
    gap_df = pd.DataFrame([{"Phrase": p, "Competitors Using It": n, "Total Uses": c} for p, n, c in gaps],
                          columns=["Phrase", "Competitors Using It", "Total Uses"])

    my_types = set((main["row"]["Schema Types"] or "").split(", ")) - {""}
    schema_gap = sorted({t for _, c in pages[1:] for t in (c["row"]["Schema Types"] or "").split(", ") if t} - my_types)
    h2s = pd.DataFrame([{"Page": label, "URL": a["url"],
                         "H2 Headings": " | ".join(h["Heading"] for h in a["details"]["headings"] if h["Level"] == "H2")}
                        for label, a in pages], columns=["Page", "URL", "H2 Headings"])
    return {"table": table, "gaps": gap_df, "schema_gap": schema_gap, "h2s": h2s}


# --------------------------------------------------------------------------
# Excel download
# --------------------------------------------------------------------------
def page_excel_bytes(audit: dict, comparison: dict = None) -> bytes:
    s = audit["score"]
    numbers = {"On-page score": s["score"], "Grade": s["grade"],
               **{f"{a} (%)": int(v) for a, v in zip(s["areas"]["Area"], s["areas"]["Score (%)"])}}
    read = audit["readability"]
    read_df = pd.DataFrame([{"Measure": k, "Value": "" if v is None else str(v)} for k, v in read.items()
                            if k not in ("examples", "status")] + [{"Measure": "Note", "Value": read.get("status", "")}])
    sheets = [("Checklist", s["checklist"]), ("Areas", s["areas"]),
              ("Headings", pd.DataFrame(audit["details"]["headings"], columns=["Level", "Heading"])),
              ("Images", pd.DataFrame(audit["details"]["images"], columns=["Image", "Alt", "Alt Text", "Width", "Height", "Loading"])),
              ("Schema fields", audit["schema"]), ("Readability", read_df),
              ("Top words", audit["content"]["words"]), ("Top phrases", audit["content"]["phrases"])]
    if audit["links"] is not None:
        sheets.append(("Links", audit["links"]))
    if audit["mobile"] is not None:
        sheets.append(("Mobile vs desktop", audit["mobile"]))
    if comparison:
        sheets += [("Competitors", comparison["table"]), ("Content gaps", comparison["gaps"]),
                   ("Competitor H2s", comparison["h2s"])]
    sheets.append(("All page data", audit["df"].drop(columns=["_issues"])))
    return checks.simple_excel("Page audit", {"URL": audit["url"], "Keyword": audit["keyword"] or "None"}, numbers, sheets)


# --------------------------------------------------------------------------
# Search result preview (approximate)
# --------------------------------------------------------------------------
DESKTOP_TITLE_PX, DESKTOP_DESC_PX, MOBILE_DESC_PX = core.TITLE_MAX_PX, 920, 680
MOBILE_TITLE_CHARS = 78


def cut_to_width(text: str, max_px: int, font_px: int) -> str:
    """The text as Google would show it: cut at a word near max_px, with an ellipsis."""
    if core.pixel_width(text, font_px) <= max_px:
        return text
    out = ""
    for word in text.split():
        trial = (out + " " + word).strip()
        if core.pixel_width(trial + " ...", font_px) > max_px:
            break
        out = trial
    return (out or text[: max(1, len(text) // 2)]) + " ..."


def serp_preview(row: dict) -> dict:
    url = row["Final URL"] or row["URL"]
    try:
        p = urlsplit(url)
        crumbs = [p.netloc] + [s for s in p.path.split("/") if s]
    except ValueError:
        crumbs = [url]
    title = row["Title"] or row["H1"] or "(no title: Google will write one)"
    desc = row["Meta Description"] or "(no meta description: Google will take text from the page)"
    mobile_title = title if len(title) <= MOBILE_TITLE_CHARS else title[: MOBILE_TITLE_CHARS - 4].rsplit(" ", 1)[0] + " ..."
    return {
        "breadcrumb": " › ".join(crumbs[:4]),
        "desktop_title": cut_to_width(title, DESKTOP_TITLE_PX, 20),
        "desktop_desc": cut_to_width(desc, DESKTOP_DESC_PX, 14),
        "mobile_title": mobile_title,
        "mobile_desc": cut_to_width(desc, MOBILE_DESC_PX, 14),
        "title_cut": core.pixel_width(row["Title"] or "", 20) > DESKTOP_TITLE_PX,
        "desc_cut": core.pixel_width(row["Meta Description"] or "", 14) > DESKTOP_DESC_PX,
    }


# --------------------------------------------------------------------------
# Word report for one page
# --------------------------------------------------------------------------
def build_page_report(audit: dict, comparison: dict = None, brand: dict = None) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    s, row = audit["score"], audit["row"]
    url = audit["url"]
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    report.apply_brand(doc, brand)
    doc.add_heading("Page audit report", level=0)
    sub = doc.add_paragraph()
    host = urlsplit(url).netloc
    sub.add_run(report.brand_line(brand) + f"Page: {report.url_label(url)} on {host}. Checked on {date.today().strftime('%d %B %Y')}"
                + (f", target keyword: {audit['keyword']}." if audit["keyword"] else ".")).italic = True
    link_p = doc.add_paragraph("Open the page: ")
    report._link(link_p, url, report.url_label(url))

    # ---- score
    doc.add_heading("On-page score", level=1)
    p = doc.add_paragraph()
    run = p.add_run(f"{s['score']} out of 100, grade {s['grade']}")
    run.bold = True
    run.font.size = Pt(16)
    counts = s["counts"]
    doc.add_paragraph(
        f"Of {sum(counts.get(k, 0) for k in ('Pass', 'Warning', 'Fail'))} scored checks, {counts.get('Pass', 0)} passed, "
        f"{counts.get('Warning', 0)} {'needs' if counts.get('Warning', 0) == 1 else 'need'} attention and "
        f"{counts.get('Fail', 0)} failed. "
        "The score is this tool's own measure of on-page SEO, not a Google metric."
    )
    if not s.get("loaded", True):
        doc.add_paragraph("The page did not load, so only indexability was checked and the score is 0: "
                          + "; ".join(s["blockers"]) + ".")
    elif s["blockers"]:
        doc.add_paragraph("These problems stop the page from ranking, so the score is capped at "
                          f"{BLOCKER_CAP}: " + "; ".join(s["blockers"]) + ".")
    report._table(doc, ["Area", "Score", "Points"],
                  [[a, f"{pct}%", f"{pts:g} of {out:g}"] for a, pct, pts, out in
                   zip(s["areas"]["Area"], s["areas"]["Score (%)"], s["areas"]["Points"], s["areas"]["Out Of"])])

    # ---- fixes, biggest point losses first
    todo = s["checklist"][s["checklist"]["Result"].isin(["Fail", "Warning"])].copy()
    todo["_lost"] = todo["Out Of"] - todo["Points"]
    todo = todo.sort_values("_lost", ascending=False)
    doc.add_heading("What to fix first", level=1)
    if todo.empty:
        doc.add_paragraph("Every scored check passed.")
    else:
        doc.add_paragraph("Ordered by the points each fix would add to the score.")
        report._table(doc, ["Check", "Found", "How to fix", "Points to gain"],
                      [[r["Check"], r["Found"], r["How To Fix"], f"{r['_lost']:g}"] for r in todo.to_dict("records")])

    # ---- search result preview (only for a page that loaded)
    if s.get("loaded", True):
        prev = serp_preview(row)
        doc.add_heading("How the page may look in Google", level=1)
        doc.add_paragraph("An estimate based on the width Google usually shows; Google may also rewrite titles and descriptions.")
        for label, t, d in (("Desktop", prev["desktop_title"], prev["desktop_desc"]),
                            ("Mobile", prev["mobile_title"], prev["mobile_desc"])):
            para = doc.add_paragraph()
            para.add_run(f"{label}: ").bold = True
            para.add_run(prev["breadcrumb"] + "\n")
            para.add_run(t + "\n").font.color.rgb = RGBColor(0x1A, 0x0D, 0xAB)
            para.add_run(d)

    # ---- readability, content, schema, devices and links (only for a page that loaded)
    if s.get("loaded", True):
        # ---- readability and content
        read = audit["readability"]
        doc.add_heading("Readability", level=1)
        if read.get("Reading ease") is not None:
            doc.add_paragraph(f"Reading ease {read['Reading ease']} out of 100 ({read['Level'].lower()}). "
                              "60 or higher suits most audiences.")
        else:
            doc.add_paragraph(read.get("status") or "Not scored.")
        report._table(doc, ["Measure", "Value"], [[k, str(read[k])] for k in
                      ("Words", "Sentences", "Average words per sentence", "Long sentences", "Long sentences (%)",
                       "Paragraphs", "Long paragraphs")])
        content = audit["content"]
        doc.add_heading("Content analysis", level=1)
        report._table(doc, ["Measure", "Value"], [[k, str(v)] for k, v in content["numbers"].items()])
        if len(content["words"]):
            doc.add_paragraph().add_run("Most used words").bold = True
            report._table(doc, ["Word", "Uses", "Density (%)"], content["words"].head(12).values.tolist())
        if len(content["phrases"]):
            doc.add_paragraph().add_run("Most used phrases").bold = True
            report._table(doc, ["Phrase", "Words", "Uses"], content["phrases"].head(12).values.tolist())

        # ---- schema
        doc.add_heading("Structured data (schema)", level=1)
        if audit["schema"].empty:
            doc.add_paragraph("No schema markup was found. " + report.advice_for("No schema markup")[1])
        else:
            report._table(doc, ["Type", "Result", "Missing required", "Missing recommended"],
                          audit["schema"][["Type", "Result", "Missing Required", "Missing Recommended"]].values.tolist())
            doc.add_paragraph("Confirm the markup with Google's Rich Results Test before and after changes.")

        # ---- smartphone vs desktop
        if audit["mobile"] is not None:
            doc.add_heading("Smartphone and desktop versions", level=1)
            diff = audit["mobile"][audit["mobile"]["Same"] == "No"]
            if diff.empty:
                doc.add_paragraph("Both versions match on every field compared.")
            else:
                doc.add_paragraph("Google indexes the smartphone version, so anything missing there is missing for Google.")
                report._table(doc, ["Field", "Smartphone", "Desktop"], diff[["Field", "Smartphone", "Desktop"]].values.tolist())
            if audit.get("mobile_note"):
                doc.add_paragraph(audit["mobile_note"])

        # ---- links
        links = audit["links"]
        if links is not None:
            bad = links[links["Result"].map(crawl.is_broken)]
            doc.add_heading("Links", level=1)
            if bad.empty:
                doc.add_paragraph(f"All {len(links):,} links on the page work.")
            else:
                doc.add_paragraph(f"{len(bad)} of {len(links):,} links on the page lead to an error.")
                report._table(doc, ["Link", "Anchor text", "Result"],
                              [[report.url_label(u, url), a, r] for u, a, r in bad[["Link URL", "Anchor Text", "Result"]].values])

    # ---- competitors
    if comparison:
        doc.add_heading("Compared with competitors", level=1)
        table = comparison["table"].copy()
        is_url = table["Measure"] == "URL"
        for col in table.columns[1:]:
            table.loc[is_url, col] = table.loc[is_url, col].map(lambda u: report.url_label(u, url) if u else u)
        table.loc[is_url, "Measure"] = "Page"
        report._table(doc, list(table.columns), table.values.tolist())
        if len(comparison["gaps"]):
            doc.add_paragraph().add_run("Phrases competitors use that this page does not").bold = True
            report._table(doc, list(comparison["gaps"].columns), comparison["gaps"].head(20).values.tolist())
        if comparison["schema_gap"]:
            doc.add_paragraph("Schema types competitors use that this page does not: " + ", ".join(comparison["schema_gap"]) + ".")

    # ---- full checklist
    doc.add_heading("Every check", level=1)
    report._table(doc, ["Area", "Check", "Result", "Found", "Points"],
                  [[r["Area"], r["Check"], r["Result"], r["Found"],
                    f"{r['Points']:g} of {r['Out Of']:g}" if r["Out Of"] else "Not scored"]
                   for r in s["checklist"].to_dict("records")])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
