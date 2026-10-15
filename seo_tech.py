"""
SEO Audit Toolkit - engine for the technical tools (no screen code in this file)

  1. Duplicate content        exact and near-duplicate main content (MinHash), repeated titles, H1s and metas
  2. Custom search            find text or code on every page, and pull out any element (CSS, XPath, regex)
  3. Staging vs live          the same pages on two sites, compared field by field, with a launch checklist
  4. Site structure           internal link score (PageRank style), folders, crawl depth, crawl tree diagram
  5. Robots.txt and AI bots   which crawlers robots.txt lets in, what the server answers them, a draft tester
"""

import hashlib
import io
import re
import zlib
from concurrent.futures import ThreadPoolExecutor
from datetime import date
from urllib.parse import urlsplit

import numpy as np
import pandas as pd
from bs4 import BeautifulSoup

import seo_checks as checks
import seo_core as core
import seo_crawl as crawl
import seo_page as page
import seo_report as report

# --------------------------------------------------------------------------
# Settings you can safely change
# --------------------------------------------------------------------------
DUP_THRESHOLD = 90           # default similarity (%) for near-duplicates
DUP_MIN_WORDS = 50           # pages with fewer main-content words are not compared (too short to judge)
DUP_SHINGLE = 3              # words per shingle
MAX_MATCHES = 50             # most values kept per extraction rule and page
LINK_SCORE_DAMPING = 0.85


def _drop(x: dict, *keys):
    for key in ("html", "headers", "main_text_marked") + keys:
        x.pop(key, None)


# ==========================================================================
# 1. Duplicate content
# ==========================================================================
_BANDS, _ROWS = 20, 6
_PERM = _BANDS * _ROWS
_RNG = np.random.RandomState(20261014)
_A = _RNG.randint(1, 2 ** 31 - 1, _PERM).astype(np.uint64)
_B = _RNG.randint(0, 2 ** 31 - 1, _PERM).astype(np.uint64)
_PRIME = np.uint64((1 << 61) - 1)


def signature(words: list):
    """MinHash signature of the text's word shingles (None when the text is too short)."""
    if len(words) < DUP_SHINGLE:
        return None
    shingles = {" ".join(words[i:i + DUP_SHINGLE]) for i in range(len(words) - DUP_SHINGLE + 1)}
    h = np.fromiter((zlib.crc32(s.encode("utf-8")) for s in shingles), dtype=np.uint64, count=len(shingles))
    return ((_A[:, None] * h[None, :] + _B[:, None]) % _PRIME).min(axis=1)


def duplicate_hook() -> dict:
    """Page-check options that keep a fingerprint of each page's main content (and drop the text)."""
    def after(row: dict):
        x = row["_x"]
        try:
            if x.get("parsed"):
                words = page.content_words(x.get("main_text", ""))
                x["dup"] = {"hash": hashlib.md5(" ".join(words).encode("utf-8")).hexdigest() if words else "",
                            "sig": signature(words) if len(words) >= DUP_MIN_WORDS else None, "words": len(words)}
        finally:
            _drop(x, "main_text", "intro_text", "h2s")
    return {"text": True, "after": after}


class _Groups:
    def __init__(self, n):
        self.parent = list(range(n))

    def find(self, i):
        while self.parent[i] != i:
            self.parent[i] = self.parent[self.parent[i]]
            i = self.parent[i]
        return i

    def join(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.parent[max(ra, rb)] = min(ra, rb)


def _main_page(members: list, rows: list) -> int:
    """The page to keep: the one the others canonicalise to, else the shortest address."""
    keys = {crawl.crawl_key(rows[i]["Final URL"] or rows[i]["URL"]): i for i in members}
    votes = {}
    for i in members:
        r = rows[i]
        if r["Canonical Status"] == "Different" and r["Canonical URL"]:
            target = crawl.crawl_key(r["Canonical URL"])
            if target in keys:
                votes[keys[target]] = votes.get(keys[target], 0) + 1
    if votes:
        return max(votes, key=votes.get)
    return min(members, key=lambda i: (rows[i]["Indexable"] != "Yes", len(rows[i]["URL"]), rows[i]["URL"]))


_BIG_BUCKET = 300   # above this many pages in one bucket, compare with a few leading pages instead of every pair
_BIG_PIVOTS = 25


def _bucket_pairs(members: list, rows: list, limit: float, seen: set, pair_sims: dict):
    """Adds the similar pairs among pages whose signatures share one band (members in row order)."""
    if len(members) <= _BIG_BUCKET:
        for a in range(len(members)):
            for b in range(a + 1, len(members)):
                i, j = members[a], members[b]
                if (i, j) in seen:
                    continue
                seen.add((i, j))
                sim = float(np.mean(rows[i]["_x"]["dup"]["sig"] == rows[j]["_x"]["dup"]["sig"]))
                if sim >= limit:
                    pair_sims[(i, j)] = sim
        return
    # Hundreds of pages built on one template (city or service pages): every pair would take too long and too
    # much memory, so each page is compared with a leading page, and pages unlike it with the next one.
    sigs = np.stack([rows[i]["_x"]["dup"]["sig"] for i in members])
    left = np.arange(len(members))
    for _ in range(_BIG_PIVOTS):
        if len(left) < 2:
            break
        p, rest = left[0], left[1:]
        sims = (sigs[rest] == sigs[p]).mean(axis=1)
        hit = sims >= limit
        for k in np.nonzero(hit)[0]:
            i, j = members[p], members[rest[k]]
            if (i, j) not in seen:
                seen.add((i, j))
                pair_sims[(i, j)] = float(sims[k])
        left = rest[~hit]


def find_duplicates(rows: list, threshold: int = DUP_THRESHOLD) -> dict:
    """Exact and near-duplicate pages among the rows checked with duplicate_hook().

    Sets Exact Duplicates, Near Duplicates, Closest Match and Similarity (%) on every row, adds the issues
    (only when two indexable pages share the content) and returns the tables for the screen and downloads."""
    limit = threshold / 100
    for r in rows:  # running again (another threshold) replaces the earlier duplicate issues
        more = r["_x"].get("more_issues")
        if more:
            r["_x"]["more_issues"] = [m for m in more if not m[1].startswith(("Exact duplicate content",
                                                                                "Near-duplicate content"))]
    ok = [i for i, r in enumerate(rows)
          if r["_x"].get("dup") and r["Final Status Code"] == 200 and not r["Redirect Hops"]]
    groups = _Groups(len(rows))
    best = {}  # row -> (similarity, other row)
    pair_sims = {}

    # exact copies: same text
    by_hash = {}
    for i in ok:
        d = rows[i]["_x"]["dup"]
        if d["hash"] and d["words"] >= 10:
            by_hash.setdefault(d["hash"], []).append(i)
    exact_of = {}
    for members in by_hash.values():
        if len(members) > 1:
            for i in members:
                exact_of[i] = members
                groups.join(members[0], i)
                best[i] = (1.0, members[1] if members[0] == i else members[0])

    # near copies: MinHash with banding, one representative per exact group
    reps = [i for i in ok if rows[i]["_x"]["dup"]["sig"] is not None and (i not in exact_of or exact_of[i][0] == i)]
    buckets = {}
    for i in reps:
        sig = rows[i]["_x"]["dup"]["sig"]
        for b in range(_BANDS):
            buckets.setdefault((b, sig[b * _ROWS:(b + 1) * _ROWS].tobytes()), []).append(i)
    seen = set()
    for members in buckets.values():
        if len(members) >= 2:
            _bucket_pairs(members, rows, limit, seen, pair_sims)

    # near matches are kept per representative page; copies of a page share its matches
    rep_of = {i: (exact_of[i][0] if i in exact_of else i) for i in ok}
    size_of = {i: len(exact_of.get(i, [i])) for i in reps}
    indexable_in = {i: sum(rows[j]["Indexable"] == "Yes" for j in exact_of.get(i, [i])) for i in set(rep_of.values())}
    near = {}
    for (i, j), sim in pair_sims.items():
        groups.join(i, j)
        for a, b in ((i, j), (j, i)):
            near.setdefault(a, {})[b] = max(near.get(a, {}).get(b, 0), sim)
            for x in exact_of.get(a, [a]):
                if x not in best or best[x][0] < sim:
                    best[x] = (sim, b)
    for members in exact_of.values():  # copies of a page share its near matches
        for i in members:
            groups.join(members[0], i)

    clusters = {}
    for i in ok:
        root = groups.find(i)
        clusters.setdefault(root, []).append(i)
    clusters = {k: v for k, v in clusters.items() if len(v) > 1}
    sims_of = {}
    for (a, b), s in pair_sims.items():
        sims_of.setdefault(groups.find(a), []).append(s)
    group_of = {}
    group_rows = []
    for n, (root, members) in enumerate(sorted(clusters.items(), key=lambda kv: -len(kv[1])), 1):
        for i in members:
            group_of[i] = n
        sims = list(sims_of.get(root, []))
        exact = all(rows[i]["_x"]["dup"]["hash"] == rows[members[0]]["_x"]["dup"]["hash"] for i in members)
        if any(i in exact_of for i in members):
            sims.append(1.0)
        indexable = [i for i in members if rows[i]["Indexable"] == "Yes"]
        main = _main_page(members, rows)
        if len(indexable) <= 1:
            action = "Handled: only one of these pages is indexable (the rest are canonicalised or noindex)."
        else:
            action = (f"{len(indexable)} indexable pages share this content. Keep {report.url_label(rows[main]['URL'])}, "
                      "then merge the others into it with a 301 redirect, point their canonical tag at it, or rewrite "
                      "them so each page is useful on its own.")
        group_rows.append({
            "Group": n, "Pages": len(members),
            "Type": "Exact copies" if exact else ("Exact and near copies" if any(i in exact_of for i in members)
                                                  else "Near copies"),
            "Indexable": len(indexable),
            "Similarity (%)": "100" if exact else f"{int(min(sims) * 100) if sims else threshold}–{int(max(sims) * 100) if sims else 100}",
            "Suggested Page To Keep": rows[main]["URL"], "Action": action,
            "URLs": " | ".join(rows[i]["URL"] for i in members[:30]) + (" ..." if len(members) > 30 else ""),
        })

    # repeated titles, H1s and meta descriptions
    repeats = []
    for col, label in (("Title", "Title"), ("H1", "H1"), ("Meta Description", "Meta description")):
        same = {}
        for i in ok:
            v = (rows[i][col] or "").strip().lower()
            if v:
                same.setdefault(v, []).append(i)
        for members in same.values():
            if len(members) > 1:
                repeats.append({"Element": label, "Text": rows[members[0]][col], "Pages": len(members),
                                "URLs": " | ".join(rows[i]["URL"] for i in members[:30])})
    rep_df = pd.DataFrame(repeats, columns=["Element", "Text", "Pages", "URLs"])
    if not rep_df.empty:
        rep_df = rep_df.sort_values(["Pages", "Element"], ascending=[False, True]).reset_index(drop=True)

    # every page
    same_counts = {}
    for col in ("Title", "H1", "Meta Description"):
        counts = {}
        for i in ok:
            v = (rows[i][col] or "").strip().lower()
            if v:
                counts[v] = counts.get(v, 0) + 1
        same_counts[col] = counts
    page_rows = []
    for i, r in enumerate(rows):
        d = r["_x"].get("dup")
        sim, other = best.get(i, (None, None))
        n_exact = len(exact_of.get(i, [i])) - 1
        matches = near.get(rep_of[i], {}) if i in rep_of else {}
        n_near = sum(size_of[q] for q in matches)
        r["Exact Duplicates"], r["Near Duplicates"] = n_exact, n_near
        r["Closest Match"] = rows[other]["URL"] if other is not None else ""
        r["Similarity (%)"] = int(round(sim * 100)) if sim is not None else None
        if i in ok and r["Indexable"] == "Yes":
            issues = r["_x"].setdefault("more_issues", [])
            exact_idx = indexable_in[rep_of[i]] - 1 if i in exact_of else 0
            near_idx = sum(indexable_in[q] for q in matches)
            if exact_idx:
                issues.append((core.WARNING, "Exact duplicate content", f"Exact duplicate content ({exact_idx} other pages)"))
            elif near_idx:
                issues.append((core.NOTICE, f"Near-duplicate content ({threshold}%+ similar)",
                               f"Near-duplicate content ({near_idx} pages, up to {int(round(sim * 100))}% similar)"))

        def same(col):
            v = (r[col] or "").strip().lower()
            return max(same_counts[col].get(v, 0) - 1, 0) if v and i in ok else 0
        page_rows.append({
            "URL": r["URL"], "Status": r["Final Status Code"], "Indexable": r["Indexable"],
            "Words": d["words"] if d else None, "Group": group_of.get(i), "Exact Duplicates": n_exact,
            "Near Duplicates": n_near, "Closest Match": r["Closest Match"], "Similarity (%)": r["Similarity (%)"],
            "Same Title As": same("Title"), "Same H1 As": same("H1"), "Same Meta As": same("Meta Description"),
            "Canonical Status": r["Canonical Status"], "Title": r["Title"],
            "Note": "" if i in ok else (r["Fetch Error"] or f"Not compared (status {r['Final Status Code']})"),
        })
    pages_df = pd.DataFrame(page_rows)
    for c in ("Status", "Words", "Group", "Similarity (%)"):
        pages_df[c] = pd.array(pages_df[c].tolist(), dtype="Int64")
    pairs = []
    for (i, j), sim in sorted(pair_sims.items(), key=lambda kv: -kv[1]):
        pairs.append({"Page": rows[i]["URL"], "Similar Page": rows[j]["URL"], "Similarity (%)": int(round(sim * 100)),
                      "Words": rows[i]["_x"]["dup"]["words"], "Similar Page Words": rows[j]["_x"]["dup"]["words"],
                      "Both Indexable": "Yes" if rows[i]["Indexable"] == rows[j]["Indexable"] == "Yes" else "No"})
    for members in {tuple(m) for m in exact_of.values()}:
        for j in members[1:]:
            pairs.append({"Page": rows[members[0]]["URL"], "Similar Page": rows[j]["URL"], "Similarity (%)": 100,
                          "Words": rows[members[0]]["_x"]["dup"]["words"], "Similar Page Words": rows[j]["_x"]["dup"]["words"],
                          "Both Indexable": "Yes" if rows[members[0]]["Indexable"] == rows[j]["Indexable"] == "Yes" else "No"})
    pairs_df = pd.DataFrame(pairs, columns=["Page", "Similar Page", "Similarity (%)", "Words", "Similar Page Words",
                                            "Both Indexable"]).sort_values("Similarity (%)", ascending=False)
    groups_df = pd.DataFrame(group_rows, columns=["Group", "Type", "Pages", "Indexable", "Similarity (%)",
                                                  "Suggested Page To Keep", "Action", "URLs"])
    compared = len([i for i in ok if rows[i]["_x"]["dup"]["sig"] is not None])
    numbers = {"Pages compared": compared,
               "Exact copies": int((pages_df["Exact Duplicates"] > 0).sum()),
               "Near copies": int(((pages_df["Near Duplicates"] > 0) & (pages_df["Exact Duplicates"] == 0)).sum()),
               "Groups to fix": int((groups_df["Indexable"] > 1).sum()) if len(groups_df) else 0,
               "Repeated titles": int((rep_df["Element"] == "Title").sum()) if len(rep_df) else 0,
               "Too short to compare": len(ok) - compared}
    return {"pages": pages_df, "groups": groups_df, "pairs": pairs_df.reset_index(drop=True), "repeats": rep_df,
            "numbers": numbers, "threshold": threshold}


def duplicates_excel(result: dict, info: dict) -> bytes:
    return checks.simple_excel("Duplicate content", info, result["numbers"],
                               [("Groups", result["groups"]), ("Similar pairs", result["pairs"]),
                                ("Repeated titles H1s metas", result["repeats"]), ("All pages", result["pages"])])


def duplicates_section(result: dict, limit: int = 25) -> dict:
    g = result["groups"]
    if g is None or g.empty:
        return {}
    todo = g[g["Indexable"] > 1].head(limit)
    if todo.empty:
        return {}
    return {"title": "Duplicate content",
            "intro": f"Groups of indexable pages with the same or nearly the same main content ({result['threshold']}% "
                     "similar or more). Google shows one page from a group of copies, so the others rarely rank.",
            "header": ["Group", "Type", "Pages", "Similarity (%)", "Page to keep", "What to do"],
            "rows": [[str(r["Group"]), r["Type"], str(r["Pages"]), r["Similarity (%)"],
                      report.url_label(r["Suggested Page To Keep"]), r["Action"]] for r in todo.to_dict("records")]}


# ==========================================================================
# 2. Custom search and extraction
# ==========================================================================
SEARCH_COLS = ["Name", "Mode", "Type", "Look in", "Find", "Match case"]
EXTRACT_COLS = ["Name", "Method", "Expression", "Extract", "Attribute"]
SEARCH_MODES = ["Contains", "Does not contain"]
SEARCH_TYPES = ["Text", "Regex"]
SEARCH_WHERE = ["Page source (HTML)", "Visible text"]
EXTRACT_METHODS = ["CSS selector", "XPath", "Regex"]
EXTRACT_WHAT = ["Text", "Inner HTML", "Whole element (HTML)", "Attribute"]

EXAMPLE_SEARCH = pd.DataFrame([
    {"Name": "No analytics tag (GA4 or GTM)", "Mode": "Does not contain", "Type": "Regex", "Look in": "Page source (HTML)",
     "Find": r"gtag/js\?id=G-|googletagmanager\.com/gtm\.js|GTM-[A-Z0-9]{4,}", "Match case": False},
    {"Name": "Lorem ipsum left in", "Mode": "Contains", "Type": "Text", "Look in": "Visible text",
     "Find": "lorem ipsum", "Match case": False},
    {"Name": "Out of stock", "Mode": "Contains", "Type": "Text", "Look in": "Visible text",
     "Find": "out of stock", "Match case": False},
], columns=SEARCH_COLS)
EXAMPLE_EXTRACT = pd.DataFrame([
    {"Name": "Author", "Method": "XPath", "Expression": "//meta[@name='author']/@content", "Extract": "Text", "Attribute": ""},
    {"Name": "Published date", "Method": "XPath", "Expression": "//meta[@property='article:published_time']/@content",
     "Extract": "Text", "Attribute": ""},
    {"Name": "Breadcrumb", "Method": "CSS selector",
     "Expression": "nav[aria-label*='readcrumb'] a, .breadcrumb a, .breadcrumbs a, .rank-math-breadcrumb a, .yoast-breadcrumbs a",
     "Extract": "Text", "Attribute": ""},
    {"Name": "Price in schema", "Method": "Regex", "Expression": r'"price"\s*:\s*"?([0-9][0-9.,]*)', "Extract": "Text",
     "Attribute": ""},
    {"Name": "og:image", "Method": "CSS selector", "Expression": "meta[property='og:image']", "Extract": "Attribute",
     "Attribute": "content"},
], columns=EXTRACT_COLS)
EMPTY_SEARCH = pd.DataFrame(columns=SEARCH_COLS).astype({"Match case": bool})
EMPTY_EXTRACT = pd.DataFrame(columns=EXTRACT_COLS)


def _cell(v) -> str:
    return "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()


def read_rules(search_df: pd.DataFrame, extract_df: pd.DataFrame):
    """The rules from the two tables, checked. Returns (search rules, extraction rules, problems)."""
    from lxml import etree
    import soupsieve
    problems, search, extract = [], [], []
    names = {"URL", "Status Code", "Title", "Note"}  # columns the results table already has

    def unique(name, fallback):
        name = name or fallback
        base, n = name, 2
        while name in names:
            name, n = f"{base} ({n})", n + 1
        names.add(name)
        return name
    for n, r in enumerate((search_df if search_df is not None else EMPTY_SEARCH).to_dict("records"), 1):
        find = _cell(r.get("Find"))
        if not find:
            continue
        rule = {"name": unique(_cell(r.get("Name")), f"Search {n}"), "mode": _cell(r.get("Mode")) or "Contains",
                "type": _cell(r.get("Type")) or "Text", "where": _cell(r.get("Look in")) or SEARCH_WHERE[0],
                "find": find, "case": bool(r.get("Match case")) if r.get("Match case") is not None and not pd.isna(r.get("Match case")) else False}
        flags = 0 if rule["case"] else re.I
        try:
            rule["rx"] = re.compile(find if rule["type"] == "Regex" else re.escape(find), flags)
        except re.error as e:
            problems.append(f"Search '{rule['name']}': the regex does not work ({e}).")
            continue
        search.append(rule)
    for n, r in enumerate((extract_df if extract_df is not None else EMPTY_EXTRACT).to_dict("records"), 1):
        expr = _cell(r.get("Expression"))
        if not expr:
            continue
        rule = {"name": unique(_cell(r.get("Name")), f"Extract {n}"), "method": _cell(r.get("Method")) or "CSS selector",
                "expr": expr, "what": _cell(r.get("Extract")) or "Text", "attr": _cell(r.get("Attribute"))}
        try:
            if rule["method"] == "XPath":
                rule["xp"] = etree.XPath(expr)
            elif rule["method"] == "Regex":
                rule["rx"] = re.compile(expr, re.I | re.S)
            else:
                soupsieve.compile(expr)
        except Exception as e:
            problems.append(f"Extraction '{rule['name']}': the {rule['method']} does not work ({e}).")
            continue
        if rule["what"] == "Attribute" and not rule["attr"] and rule["method"] != "Regex":
            problems.append(f"Extraction '{rule['name']}': enter the attribute name to extract (for example content or href).")
            continue
        extract.append(rule)
    return search, extract, problems


def _xpath_values(tree, rule) -> list:
    from lxml import etree
    out = []
    try:
        found = rule["xp"](tree)
    except Exception:
        return out
    if not isinstance(found, list):
        found = [found]
    for item in found:
        if isinstance(item, etree._Element):
            if rule["what"] == "Attribute":
                v = item.get(rule["attr"])
            elif rule["what"] == "Inner HTML":
                v = (item.text or "") + "".join(etree.tostring(c, encoding="unicode", method="html") for c in item)
            elif rule["what"] == "Whole element (HTML)":
                v = etree.tostring(item, encoding="unicode", method="html")
            else:
                v = item.text_content()
        else:
            v = str(item)
        if v is not None:
            out.append(v)
    return out


def _css_values(soup, rule) -> list:
    out = []
    for el in soup.select(rule["expr"], limit=MAX_MATCHES):
        if rule["what"] == "Attribute":
            v = el.get(rule["attr"])
            v = " ".join(v) if isinstance(v, list) else v
        elif rule["what"] == "Inner HTML":
            v = el.decode_contents()
        elif rule["what"] == "Whole element (HTML)":
            v = str(el)
        else:
            v = el.get_text(" ")
        if v is not None:
            out.append(v)
    return out


def search_page(html: bytes, search: list, extract: list) -> dict:
    """{rule name: count (search) or [values] (extraction)} for one page."""
    import lxml.html
    try:
        text = html.decode("utf-8")
    except UnicodeDecodeError:
        text = html.decode("latin-1", "replace")
    out = {}
    soup = BeautifulSoup(html, "lxml")
    tree = None
    for rule in extract:
        if rule["method"] == "Regex":
            values = []
            for m in rule["rx"].finditer(text):
                values.append(m.group(1) if m.groups() else m.group(0))
                if len(values) >= MAX_MATCHES:
                    break
        elif rule["method"] == "XPath":
            if tree is None:
                try:
                    tree = lxml.html.fromstring(html)
                except Exception:
                    tree = False
            values = _xpath_values(tree, rule) if tree is not False else []
        else:
            values = _css_values(soup, rule)
        out[rule["name"]] = [core.clean(v)[:1000] if rule["what"] in ("Text", "Attribute") else v.strip()[:2000]
                             for v in values[:MAX_MATCHES] if str(v).strip()]
    visible = None
    for rule in search:
        if rule["where"] == "Visible text":
            if visible is None:
                for t in soup(["script", "style", "noscript", "template"]):
                    t.decompose()
                visible = core.clean(soup.get_text(" "))
            hay = visible
        else:
            hay = text
        out[rule["name"]] = sum(1 for _ in rule["rx"].finditer(hay))
    return out


def search_hook(search: list, extract: list) -> dict:
    def after(row: dict):
        x = row["_x"]
        try:
            if x.get("parsed") and x.get("html"):
                x["custom"] = search_page(x["html"], search, extract)
        finally:
            _drop(x, "main_text", "intro_text", "h2s")  # links stay: a crawl follows them
    return {"keep_html": True, "after": after}


def search_tables(rows: list, search: list, extract: list):
    """(one row per page, every extracted value, a summary per rule)."""
    pages, values = [], []
    for r in rows:
        x = r["_x"]
        found = x.get("custom")
        rec = {"URL": r["URL"], "Status Code": r["Final Status Code"], "Title": r["Title"]}
        for rule in search:
            rec[rule["name"]] = found.get(rule["name"]) if found else None
        for rule in extract:
            vals = found.get(rule["name"], []) if found else []
            rec[rule["name"]] = " | ".join(vals)[:2000] if found else ""
            for n, v in enumerate(vals, 1):
                values.append({"URL": r["URL"], "Rule": rule["name"], "Match": n, "Value": v})
        rec["Note"] = "" if found is not None else (r["Fetch Error"] or f"Not read (status {r['Final Status Code']})")
        pages.append(rec)
    pages_df = pd.DataFrame(pages)
    pages_df["Status Code"] = pd.array(pages_df["Status Code"].tolist(), dtype="Int64")
    for rule in search:
        pages_df[rule["name"]] = pd.array(pages_df[rule["name"]].tolist(), dtype="Int64")
    read = [r for r in rows if r["_x"].get("custom") is not None]
    summary = []
    for rule in search:
        counts = [r["_x"]["custom"].get(rule["name"], 0) for r in read]
        matching = sum(1 for c in counts if (c == 0 if rule["mode"] == "Does not contain" else c > 0))
        summary.append({"Rule": rule["name"], "Kind": f"Search: {rule['mode'].lower()}", "Pages Matching": matching,
                        "Pages Read": len(read), "Total Matches": sum(counts)})
    for rule in extract:
        with_value = sum(1 for r in read if r["_x"]["custom"].get(rule["name"]))
        summary.append({"Rule": rule["name"], "Kind": f"Extract ({rule['method']})", "Pages Matching": with_value,
                        "Pages Read": len(read), "Total Matches": sum(len(r["_x"]["custom"].get(rule["name"], [])) for r in read)})
    return (pages_df, pd.DataFrame(values, columns=["URL", "Rule", "Match", "Value"]),
            pd.DataFrame(summary, columns=["Rule", "Kind", "Pages Matching", "Pages Read", "Total Matches"]))


def matching_pages(pages_df: pd.DataFrame, rule: dict) -> pd.DataFrame:
    col = pages_df[rule["name"]]
    if "mode" in rule:
        hit = (col == 0) if rule["mode"] == "Does not contain" else (col > 0)
        return pages_df[hit.fillna(False)]
    return pages_df[col.astype(str).str.len() > 0]


def search_excel(pages_df, values_df, summary_df, search_df, extract_df, info: dict) -> bytes:
    numbers = {"Pages checked": len(pages_df), "Rules": len(summary_df)}
    parts = [d.assign(Kind=k) for d, k in ((search_df, "Search"), (extract_df, "Extract")) if d is not None and len(d)]
    rules = pd.concat(parts, ignore_index=True) if parts else pd.DataFrame(columns=["Kind"])
    return checks.simple_excel("Custom search and extraction", info, numbers,
                               [("Rules summary", summary_df), ("Pages", pages_df), ("Extracted values", values_df),
                                ("Rules used", rules.fillna(""))])


# ==========================================================================
# 3. Staging vs live
# ==========================================================================
def parse_replacements(text: str) -> list:
    """'find => replace' lines -> [(find, replace), ...]."""
    out = []
    for line in (text or "").splitlines():
        if "=>" in line:
            a, b = line.split("=>", 1)
            if a.strip():
                out.append((a.strip(), b.strip()))
    return out


def base_of(url: str) -> str:
    """The address a site's pages start with: scheme, host and any folder, without the trailing slash."""
    url = core.normalize_url(url) if core.looks_like_url(url) else ""
    return url.rstrip("/") if url else ""


def map_url(url: str, from_base: str, to_base: str, replacements: list = None) -> str:
    """The same page on the other site: swap the start of the address, then apply the replacements."""
    try:
        p, b = urlsplit(url), urlsplit(from_base)
        if core.host_key(url) == core.host_key(from_base) and p.path.startswith(b.path or ""):
            rest = url[url.index(p.netloc) + len(p.netloc) + len(b.path or ""):]
            url = to_base + (rest if rest.startswith(("/", "?")) or not rest else "/" + rest)
    except ValueError:
        return url
    for find, repl in replacements or []:
        url = url.replace(find, repl)
    return url


def staging_hook(staging_host: str) -> dict:
    """Counts how often the staging address appears in each staging page's HTML (absolute links, canonical, images)."""
    host = staging_host.lower().encode()

    def after(row: dict):
        x = row["_x"]
        try:
            row["_staging_mentions"] = x["html"].lower().count(host) if x.get("html") else 0
        finally:
            _drop(x)
    return {"keep_html": True, "after": after}


def _n(v):
    return 0 if v is None or v is pd.NA else v


def _words_drop(a, b) -> bool:
    a, b = _n(a), _n(b)
    return a >= 50 and b < a * 0.7 and a - b >= 50


def compare_pair(live: dict, stage: dict, live_base: str, stage_base: str, back: list, wide_noindex: bool) -> list:
    """[(field, live value, staging value, severity, note)] for one page."""
    out = []

    def add(field, lv, sv, sev, note=""):
        out.append((field, "" if lv is None else str(lv), "" if sv is None else str(sv), sev, note))
    lfs, sfs = live["Final Status Code"], stage["Final Status Code"]
    if lfs is None or lfs >= 400:  # the live page is broken already: nothing to keep
        return out
    if sfs is None or sfs >= 400:
        add("Status", lfs, sfs if sfs is not None else stage["Fetch Error"], core.CRITICAL,
            "The page is missing on staging. Add it, or plan a 301 redirect to its new address.")
        return out
    if lfs != sfs:
        add("Status", lfs, sfs, core.WARNING)
    if stage["Redirect Hops"] and not live["Redirect Hops"]:
        add("Redirect", "None", map_url(stage["Final URL"], stage_base, live_base, back), core.WARNING,
            "The page redirects on staging.")
    if not (stage["_x"].get("parsed") and live["_x"].get("parsed")):
        return out
    if live["Indexable"] != stage["Indexable"]:
        noindex_only = "Noindex" in (stage["Indexability Reason"] or "")
        if not (noindex_only and wide_noindex):  # a staging-wide noindex is covered once, in the launch checklist
            add("Indexable", f"{live['Indexable']} {live['Indexability Reason']}".strip(),
                f"{stage['Indexable']} {stage['Indexability Reason']}".strip(), core.CRITICAL)
    elif (live["Meta Robots"] or "") != (stage["Meta Robots"] or "") and not wide_noindex:
        add("Meta robots", live["Meta Robots"], stage["Meta Robots"], core.WARNING)

    def canon(row, base_from):
        c = row["Canonical URL"]
        return crawl.crawl_key(map_url(c, base_from, live_base, back if base_from == stage_base else None)) if c else ""
    lc, sc = canon(live, live_base), canon(stage, stage_base)
    if lc != sc:
        if live["Canonical Status"] == "Self" and stage["Canonical Status"] == "Different":
            add("Canonical", live["Canonical URL"], stage["Canonical URL"], core.CRITICAL,
                "On staging the canonical tag points to another page.")
        elif not sc:
            add("Canonical", live["Canonical URL"], "Missing", core.WARNING)
        else:
            add("Canonical", live["Canonical URL"], stage["Canonical URL"], core.NOTICE)
    lt, stt = live["Title"] or "", stage["Title"] or ""
    if lt and not stt:
        add("Title", lt, "Missing", core.CRITICAL)
    elif lt.strip().lower() != stt.strip().lower():
        add("Title", lt, stt, core.WARNING, "Check the new title still has the page's target keyword.")
    lm, sm = live["Meta Description"] or "", stage["Meta Description"] or ""
    if lm and not sm:
        add("Meta description", lm, "Missing", core.WARNING)
    elif lm.strip().lower() != sm.strip().lower():
        add("Meta description", lm, sm, core.NOTICE)
    lh, sh = live["H1"] or "", stage["H1"] or ""
    if lh and not sh:
        add("H1", lh, "Missing", core.WARNING)
    elif lh.strip().lower() != sh.strip().lower():
        add("H1", lh, sh, core.NOTICE)
    if _n(stage["H1 Count"]) > 1 and _n(live["H1 Count"]) <= 1:
        add("H1 count", live["H1 Count"], stage["H1 Count"], core.NOTICE)
    if _words_drop(live["Word Count"], stage["Word Count"]):
        add("Words (main content)", live["Word Count"], stage["Word Count"], core.WARNING,
            "Much less text on staging. Check no content was lost in the move.")
    if _n(live["H2 Count"]) and abs(_n(live["H2 Count"]) - _n(stage["H2 Count"])) > max(1, _n(live["H2 Count"]) // 2):
        add("H2 headings", live["H2 Count"], stage["H2 Count"], core.NOTICE)
    lt_ = set((live["Schema Types"] or "").split(", ")) - {""}
    st_ = set((stage["Schema Types"] or "").split(", ")) - {""}
    if lt_ - st_:
        add("Schema removed", ", ".join(sorted(lt_ - st_)), "", core.WARNING)
    if st_ - lt_:
        add("Schema added", "", ", ".join(sorted(st_ - lt_)), core.NOTICE)
    if _n(live["Hreflang Count"]) and not _n(stage["Hreflang Count"]):
        add("Hreflang", live["Hreflang Count"], 0, core.WARNING, "Hreflang tags are missing on staging.")
    elif _n(live["Hreflang Count"]) != _n(stage["Hreflang Count"]):
        add("Hreflang", live["Hreflang Count"], stage["Hreflang Count"], core.NOTICE)
    li, si = _n(live["Internal Links"]), _n(stage["Internal Links"])
    if li >= 10 and si < li * 0.7:
        add("Internal links", li, si, core.NOTICE)
    if _n(stage["Alt Missing"]) > _n(live["Alt Missing"]):
        add("Images missing alt", live["Alt Missing"], stage["Alt Missing"], core.NOTICE)
    for col, sev in (("Lang", core.WARNING), ("Viewport", core.WARNING), ("OG Title", core.NOTICE)):
        if live[col] and not stage[col]:
            add(col, live[col], "Missing", sev)
    return out


def blocks_all(robots_text: str) -> bool:
    """True when robots.txt blocks Googlebot from the whole site."""
    return bool(robots_text) and not core.RobotsRules(robots_text).allowed("http://x/", "googlebot")


def compare_sites(live_rows: list, stage_rows: list, live_base: str, stage_base: str, back: list,
                  robots: dict = None) -> dict:
    """Pages, every change and the launch checklist for live rows matched one to one with staging rows.

    Check the staging pages without robots.txt (a staging robots.txt usually blocks everything); pass both
    files as robots={"Live": text, "Staging": text} to have them compared in the checklist."""
    parsed = [s for s in stage_rows if s["_x"].get("parsed")]
    noindex = [s for s in parsed if s["_x"].get("noindex_meta") or s["_x"].get("noindex_header")]
    wide = len(parsed) > 0 and len(noindex) >= 0.5 * len(parsed)
    pages, changes = [], []
    rank = {core.CRITICAL: 0, core.WARNING: 1, core.NOTICE: 2}
    for live, stage in zip(live_rows, stage_rows):
        diffs = compare_pair(live, stage, live_base, stage_base, back, wide)
        sfs, lfs = stage["Final Status Code"], live["Final Status Code"]
        if lfs is None or lfs >= 400:
            result = f"Live page answers {lfs if lfs is not None else live['Fetch Error']}"
        elif sfs is None or sfs >= 400:
            result = f"Missing on staging ({sfs if sfs is not None else stage['Fetch Error']})"
        elif stage["Redirect Hops"] and not live["Redirect Hops"]:
            result = "Redirects on staging"
        elif live["Final Status Code"] != 200:
            result = f"Live page answers {live['Final Status Code'] or live['Fetch Error']}"
        else:
            result = "Changed" if diffs else "Same"
        worst = min((d[3] for d in diffs), key=rank.get, default="OK")
        pages.append({"Live URL": live["URL"], "Staging URL": stage["URL"], "Result": result, "Severity": worst,
                      "Changes": len(diffs), "Changed Fields": ", ".join(d[0] for d in diffs),
                      "Live Status": live["Final Status Code"], "Staging Status": sfs,
                      "Live Title": live["Title"], "Staging Title": stage["Title"],
                      "Staging Address In HTML": stage.get("_staging_mentions") or 0})
        for field, lv, sv, sev, note in diffs:
            changes.append({"Live URL": live["URL"], "Field": field, "Live": lv, "Staging": sv, "Severity": sev,
                            "Note": note})
    pages_df = pd.DataFrame(pages)
    for c in ("Live Status", "Staging Status", "Changes", "Staging Address In HTML"):
        pages_df[c] = pd.array(pages_df[c].tolist(), dtype="Int64")
    changes_df = pd.DataFrame(changes, columns=["Live URL", "Field", "Live", "Staging", "Severity", "Note"])
    if not changes_df.empty:
        changes_df = changes_df.assign(_r=changes_df["Severity"].map(rank)).sort_values(["_r", "Field"]).drop(columns="_r")

    # launch checklist
    checklist = []
    missing = pages_df[pages_df["Result"].str.startswith("Missing")]
    stage_codes = [s["Final Status Code"] for s in stage_rows]
    if stage_rows and all(c == 401 for c in stage_codes):
        checklist.append((core.CRITICAL, "Staging asks for a login",
                          "Every staging page answered 401. Enter the staging user name and password, then compare again."))
    if len(missing):
        checklist.append((core.CRITICAL, f"{len(missing)} live page{'s are' if len(missing) != 1 else ' is'} missing on staging",
                          "Add these pages to the new site, or 301 redirect each one to its closest new page at launch. "
                          "The Redirect check tool can test the redirects after launch."))
    robots = robots or {}
    if blocks_all(robots.get("Staging", "")):
        live_blocked = blocks_all(robots.get("Live", ""))
        checklist.append((core.CRITICAL if live_blocked else core.WARNING,
                          "Staging robots.txt blocks all crawlers" + (" and so does the live one" if live_blocked else ""),
                          "Normal for a staging site, but the live site must not go live with it: keep the live "
                          "robots.txt (with its Sitemap line) at launch."))
    if wide:
        checklist.append((core.WARNING, f"{len(noindex)} of {len(parsed)} staging pages are noindex",
                          "Normal for a staging site. Remove the noindex (meta robots and X-Robots-Tag) when the site goes live."))
    mentions = pages_df[pages_df["Staging Address In HTML"].fillna(0) > 0]
    if len(mentions):
        host = urlsplit(stage_base).netloc
        checklist.append((core.WARNING, f"{len(mentions)} staging page{'s' if len(mentions) != 1 else ''} "
                                        f"mention{'s' if len(mentions) == 1 else ''} {host} in the HTML",
                          "Absolute links, canonical tags, images or scripts point at the staging address. Change them to "
                          "the live address (or relative links) before launch, or the live site will link to staging."))
    keyword_advice = ("Review them in the changes table. Titles and headings that lose the target keyword are the "
                      "usual cause of ranking drops after a redesign.")
    advice = {"Canonical": "Each page should keep a canonical tag pointing to its own live address (not the staging "
                           "address, and not another page).",
              "Words (main content)": "Check no content was lost in the move: text that ranked on the old page should "
                                      "still be on the new one.",
              "Schema removed": "Add the same structured data to the new templates (Organization, Product, Article and "
                                "so on), then test them with Google's Rich Results Test.",
              "Hreflang": "Add the hreflang tags to the new templates, with every language version linking back."}
    for field, label in (("Canonical", "canonical tags"), ("Title", "titles"), ("Meta description", "meta descriptions"),
                         ("H1", "H1 headings"), ("Words (main content)", "pages with much less text"),
                         ("Schema removed", "pages that lose schema types"), ("Hreflang", "hreflang differences")):
        n = int((changes_df["Field"] == field).sum()) if len(changes_df) else 0
        if n:
            sev = changes_df[changes_df["Field"] == field]["Severity"].map(rank).min()
            sev = [core.CRITICAL, core.WARNING, core.NOTICE][int(sev)]
            checklist.append((sev, f"{n} change{'s' if n != 1 else ''} in {label}", advice.get(field, keyword_advice)))
    checklist.sort(key=lambda c: rank[c[0]])
    numbers = {"Pages compared": len(pages_df), "Same": int((pages_df["Result"] == "Same").sum()),
               "Changed": int((pages_df["Result"] == "Changed").sum()), "Missing on staging": len(missing),
               "Critical changes": int((changes_df["Severity"] == core.CRITICAL).sum()) if len(changes_df) else 0}
    return {"pages": pages_df, "changes": changes_df,
            "checklist": pd.DataFrame(checklist, columns=["Severity", "Check", "What to do"]), "numbers": numbers,
            "noindex_wide": wide, "robots": {k: v for k, v in robots.items() if v}}


def staging_excel(result: dict, info: dict) -> bytes:
    sheets = [("Launch checklist", result["checklist"]), ("Pages", result["pages"]), ("All changes", result["changes"]),
              ("Missing on staging", result["pages"][result["pages"]["Result"].str.startswith("Missing")])]
    if result.get("new") is not None:
        sheets.append(("Only on staging", result["new"]))
    if result.get("robots"):
        sheets.append(("robots.txt", pd.DataFrame([{"Site": k, "robots.txt": v} for k, v in result["robots"].items()])))
    return checks.simple_excel("Staging vs live", info, result["numbers"], sheets)


def staging_report(result: dict, info: dict, brand: dict = None) -> bytes:
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(10.5)
    report.apply_brand(doc, brand)
    doc.add_heading("Launch check: staging vs live", level=0)
    def host(u):
        try:
            p = urlsplit(u)
            return p.netloc + (p.path.rstrip("/") if p.path not in ("", "/") else "")
        except ValueError:
            return u
    doc.add_paragraph().add_run(report.brand_line(brand) + f"Live site: {host(info.get('Live site', ''))}. Staging site: "
                                f"{host(info.get('Staging site', ''))}. Checked on {date.today().strftime('%d %B %Y')}.").italic = True
    doc.add_heading("Summary", level=1)
    report._table(doc, ["Measure", "Pages"], [[k, f"{v:,}"] for k, v in result["numbers"].items()])
    doc.add_heading("Launch checklist", level=1)
    if result["checklist"].empty:
        doc.add_paragraph("No launch problems were found.")
    else:
        report._table(doc, ["Priority", "Check", "What to do"], result["checklist"].values.tolist())
    ch = result["changes"]
    important = ch[ch["Severity"].isin([core.CRITICAL, core.WARNING])].head(150) if len(ch) else ch
    doc.add_heading("Changes to review", level=1)
    if important.empty:
        doc.add_paragraph("No critical changes or warnings.")
    else:
        doc.add_paragraph("Critical changes first, then warnings. The full list is in the Excel download.")
        report._table(doc, ["Page", "Field", "Live", "Staging", "Priority"],
                      [[report.url_label(r["Live URL"]), r["Field"], r["Live"][:120], r["Staging"][:120], r["Severity"]]
                       for r in important.to_dict("records")])
    missing = result["pages"][result["pages"]["Result"].str.startswith("Missing")]
    if len(missing):
        doc.add_heading("Live pages missing on staging", level=1)
        doc.add_paragraph("Each needs to exist on the new site or get a 301 redirect at launch.")
        report._table(doc, ["Live page", "Staging answer"],
                      [[report.url_label(r["Live URL"]), r["Result"]] for r in missing.head(300).to_dict("records")])
    new = result.get("new")
    if new is not None and len(new):
        doc.add_heading("Pages only on staging", level=1)
        doc.add_paragraph("New pages. Add them to the XML sitemap and internal links at launch.")
        report._table(doc, ["Staging page", "Live address after launch"],
                      [[report.url_label(r["URL"]), report.url_label(r["Live Address After Launch"])]
                       for r in new.head(200).to_dict("records")])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ==========================================================================
# 4. Site structure and internal link score
# ==========================================================================
def link_scores(rows: list, links: list) -> dict:
    """Internal link score (0-100) per crawled URL, PageRank style: links from strong pages count for more.
    Follows redirects to the page they land on and ignores nofollow links. 0-100 on a log scale like
    Screaming Frog's Link Score."""
    keys = [crawl.crawl_key(r["URL"]) for r in rows]
    index = {k: i for i, k in enumerate(keys)}
    final = {}
    for i, r in enumerate(rows):
        if r["Redirect Hops"] and r["Final URL"]:
            j = index.get(crawl.crawl_key(r["Final URL"]))
            if j is not None:
                final[i] = j
    edges = set()
    for source, target, _, nofollow, internal in links:
        if not internal or nofollow:
            continue
        a, b = index.get(crawl.crawl_key(source)), index.get(crawl.crawl_key(target))
        if a is None or b is None:
            continue
        b = final.get(b, b)
        if a != b:
            edges.add((a, b))
    n = len(rows)
    if not n:
        return {}
    src = np.array([e[0] for e in edges], dtype=np.int64)
    dst = np.array([e[1] for e in edges], dtype=np.int64)
    out_deg = np.bincount(src, minlength=n).astype(float) if len(edges) else np.zeros(n)
    pr = np.full(n, 1.0 / n)
    d = LINK_SCORE_DAMPING
    for _ in range(50):
        new = np.full(n, (1 - d) / n)
        if len(edges):
            np.add.at(new, dst, d * pr[src] / out_deg[src])
        new += d * pr[out_deg == 0].sum() / n
        if np.abs(new - pr).sum() < 1e-9:
            pr = new
            break
        pr = new
    lo, hi = np.log(pr.min()), np.log(pr.max())
    scores = np.zeros(n) if hi - lo < 1e-12 else 100 * (np.log(pr) - lo) / (hi - lo)
    return {keys[i]: int(round(scores[i])) for i in range(n)}


def apply_link_scores(rows: list, links: list):
    scores = link_scores(rows, links)
    for r in rows:
        r["Link Score"] = scores.get(crawl.crawl_key(r["URL"]))


def folder_of(url: str, depth: int = 1) -> str:
    try:
        parts = [p for p in urlsplit(url).path.split("/") if p]
    except ValueError:
        return "Not a valid URL"
    if not parts:
        return "/"
    if len(parts) == 1 and "." not in parts[0] and not url.endswith("/"):
        return "/ (top-level pages)"
    folders = parts[:-1] if not url.endswith("/") else parts
    return "/" + "/".join(folders[:depth]) + "/" if folders else "/ (top-level pages)"


def structure_tables(df: pd.DataFrame) -> dict:
    """Folders, crawl depth and the pages that look buried."""
    work = df.assign(_folder=[folder_of(u) for u in df["URL"]])
    folders = []
    for name, g in work.groupby("_folder", sort=False):
        depth = g["Crawl Depth"].dropna()
        ls = g["Link Score"].dropna() if "Link Score" in g.columns else pd.Series(dtype=float)
        rec = {"Folder": name, "Pages": len(g), "Indexable": int((g["Indexable"] == "Yes").sum()),
               "Average Depth": round(float(depth.mean()), 1) if len(depth) else None,
               "Deepest": int(depth.max()) if len(depth) else None,
               "Average Link Score": int(round(float(ls.mean()))) if len(ls) else None,
               "With Critical Issues": int((g["Severity"] == core.CRITICAL).sum()),
               "With Warnings": int((g["Severity"] == core.WARNING).sum())}
        if "On-page Score" in g.columns:
            sc = g["On-page Score"].dropna()
            rec["Average On-page Score"] = int(round(float(sc.mean()))) if len(sc) else None
        folders.append(rec)
    folders_df = pd.DataFrame(folders).sort_values("Pages", ascending=False).reset_index(drop=True)
    for c in ("Deepest", "Average Link Score", "Average On-page Score"):
        if c in folders_df.columns:
            folders_df[c] = pd.array(folders_df[c].tolist(), dtype="Int64")
    crawled = df[df["Crawl Depth"].notna()]
    depth_df = (crawled.assign(_ix=crawled["Indexable"] == "Yes")
                .groupby("Crawl Depth")["_ix"].agg(["sum", "count"]).reset_index())
    depth_df = pd.DataFrame({"Clicks From Start": depth_df["Crawl Depth"].astype(int),
                             "Indexable": depth_df["sum"].astype(int),
                             "Not Indexable": (depth_df["count"] - depth_df["sum"]).astype(int)})
    ok = df[(df["Indexable"] == "Yes") & (df["Final Status Code"] == 200)]
    buried = ok.sort_values(["Link Score", "Crawl Depth"], ascending=[True, False], na_position="first").head(25)
    buried = buried[["URL", "Link Score", "Crawl Depth", "Inlinks", "Title"]]
    return {"folders": folders_df, "depth": depth_df, "buried": buried.reset_index(drop=True)}


def _short(path: str, n: int = 34) -> str:
    return path if len(path) <= n else "…" + path[-(n - 1):]


TREE_COLORS = {"ok": ("#e3f4e3", "#0ca30c"), "redirect": ("#fff3d6", "#c98a00"), "error": ("#fbe1e1", "#d03b3b"),
               "noindex": ("#ececec", "#898781")}


def crawl_tree_dot(df: pd.DataFrame, limit: int = 120, max_depth: int = 3) -> str:
    """Graphviz DOT of the crawl tree: each page under the page it was first found on."""
    rows = df[df["Crawl Depth"].notna() & (df["Crawl Depth"] <= max_depth)].sort_values("Crawl Depth").head(limit)
    keys = {crawl.crawl_key(u): n for n, u in enumerate(rows["URL"])}
    lines = ['digraph G {', 'graph [rankdir=LR, bgcolor="transparent", nodesep=0.08, ranksep=0.5];',
             'node [shape=box, style="rounded,filled", fontname="Helvetica", fontsize=10, height=0.25, margin="0.08,0.03"];',
             'edge [color="#9a9a9a", arrowsize=0.5];']
    for n, r in enumerate(rows.to_dict("records")):
        fs, hops = r["Final Status Code"], r["Redirect Hops"]
        if r.get("Robots.txt") == "Blocked":
            kind = "noindex"  # not crawled: blocked by robots.txt
        else:
            kind = ("error" if (fs is None or pd.isna(fs) or fs >= 400) else "redirect" if (hops and not pd.isna(hops) and hops > 0)
                    else "ok" if r["Indexable"] == "Yes" else "noindex")
        fill, line = TREE_COLORS[kind]
        try:
            path = urlsplit(r["URL"]).path or "/"
        except ValueError:
            path = r["URL"]
        label = _short(path).replace('"', "'")
        lines.append(f'n{n} [label="{label}", fillcolor="{fill}", color="{line}", fontcolor="#1f1f1f", tooltip="{r["URL"]}"];')
        parent = keys.get(crawl.crawl_key(r["Linked From"])) if r["Linked From"] else None
        if parent is not None and parent != n:
            lines.append(f"n{parent} -> n{n};")
    lines.append("}")
    return "\n".join(lines)


def structure_section(tables: dict) -> dict:
    f = tables["folders"]
    if f is None or f.empty:
        return {}
    cols = [c for c in ["Folder", "Pages", "Indexable", "Average Depth", "Average Link Score", "With Critical Issues"]
            if c in f.columns]
    return {"title": "Site structure",
            "intro": "Pages by folder, how many clicks they are from the start page, and their internal link score "
                     "(0 to 100: how much internal linking points at them). Important pages deep in the site or with "
                     "a low link score need more internal links.",
            "header": cols, "rows": [["" if pd.isna(v) else str(v) for v in r] for r in f[cols].head(40).values.tolist()]}


# ==========================================================================
# 5. Robots.txt and AI crawlers
# ==========================================================================
# (token, owner, what it does, kind, User-Agent string or None for tokens that do not crawl)
BOTS = [
    ("Googlebot", "Google", "Google Search", "Search engine", core.USER_AGENTS["Googlebot Smartphone"]),
    ("Bingbot", "Microsoft", "Bing search, also used by Copilot", "Search engine", core.USER_AGENTS["Bingbot"]),
    ("OAI-SearchBot", "OpenAI", "Shows sites in ChatGPT search results", "AI search", core.USER_AGENTS["OAI-SearchBot (ChatGPT search)"]),
    ("ChatGPT-User", "OpenAI", "Opens a page when a ChatGPT user asks", "User request",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko); compatible; ChatGPT-User/1.0; +https://openai.com/bot"),
    ("GPTBot", "OpenAI", "Collects pages to train OpenAI models", "AI training", core.USER_AGENTS["GPTBot (OpenAI)"]),
    ("Claude-SearchBot", "Anthropic", "Shows sites in Claude's search results", "AI search",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko; compatible; Claude-SearchBot/1.0; +https://www.anthropic.com)"),
    ("Claude-User", "Anthropic", "Opens a page when a Claude user asks", "User request",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko; compatible; Claude-User/1.0; +Claude-User@anthropic.com)"),
    ("ClaudeBot", "Anthropic", "Collects pages to train Anthropic models", "AI training", core.USER_AGENTS["ClaudeBot (Anthropic)"]),
    ("PerplexityBot", "Perplexity", "Shows sites in Perplexity answers", "AI search", core.USER_AGENTS["PerplexityBot"]),
    ("Perplexity-User", "Perplexity", "Opens a page when a Perplexity user asks", "User request",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko; compatible; Perplexity-User/1.0; +https://perplexity.ai/perplexity-user)"),
    ("Google-Extended", "Google", "Use of the site in Gemini and Google AI training (a robots.txt token; Googlebot does the crawling)",
     "AI training", None),
    ("Applebot-Extended", "Apple", "Use of the site to train Apple's AI models (a robots.txt token only)", "AI training", None),
    ("Meta-ExternalAgent", "Meta", "Collects pages to train Meta AI", "AI training",
     "meta-externalagent/1.1 (+https://developers.facebook.com/docs/sharing/webmasters/crawler)"),
    ("Amazonbot", "Amazon", "Alexa answers and Amazon AI", "AI search",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko; compatible; Amazonbot/0.1; +https://developer.amazon.com/support/amazonbot) Chrome/119.0.6045.214 Safari/537.36"),
    ("DuckAssistBot", "DuckDuckGo", "DuckDuckGo AI answers", "AI search",
     "DuckAssistBot/1.2; (+http://duckduckgo.com/duckassistbot.html)"),
    ("MistralAI-User", "Mistral AI", "Opens a page when a Le Chat user asks", "User request",
     "Mozilla/5.0 AppleWebKit/537.36 (KHTML, like Gecko; compatible; MistralAI-User/1.0; +https://docs.mistral.ai/robots)"),
    ("CCBot", "Common Crawl", "Open web archive many AI models are trained on", "AI training", "CCBot/2.0 (https://commoncrawl.org/faq/)"),
    ("Bytespider", "ByteDance", "Collects pages to train ByteDance AI", "AI training",
     "Mozilla/5.0 (Linux; Android 5.0) AppleWebKit/537.36 (KHTML, like Gecko) Mobile Safari/537.36 (compatible; Bytespider; https://zhanzhang.toutiao.com/)"),
]
BOT_KINDS = ["Search engine", "AI search", "User request", "AI training"]
GOOGLE_SUPPORTED = {"user-agent", "allow", "disallow", "sitemap"}
ROBOTS_MAX_BYTES = 500 * 1024


def robots_notes(text: str, size: int) -> list:
    """Things in a robots.txt file worth knowing: rules Google ignores, size, groups."""
    notes, unknown = [], {}
    for line in (text or "").splitlines():
        line = line.split("#", 1)[0].strip()
        if ":" in line:
            key = line.split(":", 1)[0].strip().lower()
            if key not in GOOGLE_SUPPORTED:
                unknown[key] = unknown.get(key, 0) + 1
    if "noindex" in unknown:
        notes.append("Has Noindex lines. Google stopped supporting noindex in robots.txt in 2019; use a meta robots tag.")
    if "crawl-delay" in unknown:
        notes.append("Has Crawl-delay. Google ignores it; Bing and some other crawlers follow it.")
    others = [k for k in unknown if k not in ("noindex", "crawl-delay")]
    if others:
        notes.append("Lines Google does not use: " + ", ".join(sorted(others)) + ".")
    if size > ROBOTS_MAX_BYTES:
        notes.append(f"The file is {size // 1024} KB. Google reads only the first 500 KB.")
    return notes


def bot_rules_table(rules, base: str, paths: list) -> pd.DataFrame:
    """For every crawler: which group of rules it follows, and Allowed / Blocked for each test address."""
    out = []
    for token, owner, purpose, kind, _ in BOTS:
        rec = {"Crawler": token, "Owner": owner, "Type": kind, "What It Does": purpose}
        if rules is None:
            rec["Rules Followed"] = "No robots.txt (everything allowed)"
        else:
            rec["Rules Followed"] = rules.group_for(token.lower())
        for path in paths:
            if rules is None:
                rec[path] = "Allowed"
                continue
            ok, rule = rules.match(base + path, token.lower())
            rec[path] = ("Allowed" if ok else "Blocked") + (f" ({rule})" if rule else "")
        out.append(rec)
    return pd.DataFrame(out)


def server_answers(base: str, timeout: int, workers: int = 8, headers: dict = None) -> dict:
    """What the homepage answers to each crawler's User-Agent, and to Chrome."""
    agents = {token: ua for token, _, _, _, ua in BOTS if ua}
    agents["Chrome (for comparison)"] = core.USER_AGENTS["Chrome (desktop)"]

    def one(item):
        name, ua = item
        res = core.fetch_raw(core.make_session(1, ua, headers), base + "/", timeout, max_bytes=20_000, retries=0)
        return name, res["status"], res["error"]
    with ThreadPoolExecutor(max_workers=workers) as ex:
        return {name: (status, error) for name, status, error in ex.map(one, agents.items())}


def answer_label(status, error, chrome_status) -> str:
    if status is None:
        return f"No answer ({error or 'failed'})"
    if status == 200:
        return "200 OK"
    if chrome_status == 200 and status in (401, 403, 406, 429, 503):
        return f"{status}: blocked by the server or firewall"
    return str(status)


def check_ai_access(base: str, timeout: int, paths: list, headers: dict = None, auth=None) -> dict:
    """robots.txt for every crawler, the server's answer to each, and llms.txt."""
    sess = core.make_session(4, core.USER_AGENTS["Chrome (desktop)"], headers, auth)
    rb = core.fetch_raw(sess, base + "/robots.txt", timeout, max_bytes=600_000)
    text = rb["body"].decode("utf-8", "replace") if rb["status"] == 200 else ""
    rules = core.RobotsRules(text) if rb["status"] == 200 else None
    answers = server_answers(base, timeout, headers=headers)
    chrome = answers.get("Chrome (for comparison)", (None, ""))[0]
    table = bot_rules_table(rules, base, paths)
    table.insert(5, "Server Answer (homepage)", [
        answer_label(*answers[token], chrome) if token in answers else "Does not crawl (token only)"
        for token in table["Crawler"]])
    llms = core.fetch_raw(sess, base + "/llms.txt", timeout, max_bytes=200_000)
    llms_text = llms["body"].decode("utf-8", "replace") if llms["status"] == 200 else ""
    is_llms = llms["status"] == 200 and "<html" not in llms_text[:500].lower() and bool(llms_text.strip())
    if rb["status"] == 200:
        robots_state = "Found"
    elif rb["status"] is None or rb["status"] >= 500 or rb["status"] == 429:
        robots_state = f"Could not be read ({rb['status'] or rb['error']})"
    else:
        robots_state = f"Not found ({rb['status']}): every crawler may crawl everything"
    return {"base": base, "robots_status": rb["status"], "robots_state": robots_state, "robots_text": text,
            "robots_size": len(rb["body"]), "rules": rules, "table": table, "chrome": chrome,
            "notes": robots_notes(text, len(rb["body"])), "sitemaps": rules.sitemaps if rules else [],
            "llms": {"found": is_llms, "status": llms["status"], "first_line": llms_text.strip().splitlines()[0][:120] if is_llms else ""}}


def draft_compare(result: dict, draft_text: str, paths: list) -> pd.DataFrame:
    """The live rules next to a draft robots.txt: one row per crawler and address that changes."""
    draft = core.RobotsRules(draft_text)
    live = result["rules"]
    out = []
    for token, owner, _, kind, _ in BOTS:
        for path in paths:
            url = result["base"] + path
            lv = live.match(url, token.lower()) if live else (True, "")
            dv = draft.match(url, token.lower())
            out.append({"Crawler": token, "Type": kind, "Address": path,
                        "Live": ("Allowed" if lv[0] else "Blocked") + (f" ({lv[1]})" if lv[1] else ""),
                        "Draft": ("Allowed" if dv[0] else "Blocked") + (f" ({dv[1]})" if dv[1] else ""),
                        "Changes": "Yes" if lv[0] != dv[0] else "No"})
    return pd.DataFrame(out)


def ai_summary(table: pd.DataFrame, path: str = "/") -> dict:
    blocked = table[path].str.startswith("Blocked")
    firewall = table["Server Answer (homepage)"].str.contains("blocked by the server", na=False)
    out = {}
    for kind in BOT_KINDS:
        part = table["Type"] == kind
        out[f"{kind}: allowed"] = f"{int((part & ~blocked).sum())} of {int(part.sum())}"
    out["Blocked by server"] = int(firewall.sum())
    return out


def ai_excel(result: dict, draft: pd.DataFrame, info: dict) -> bytes:
    numbers = {"robots.txt": result["robots_state"], "llms.txt": "Found" if result["llms"]["found"] else "Not found",
               **ai_summary(result["table"])}
    sheets = [("Crawlers", result["table"]),
              ("Notes", pd.DataFrame({"Note": result["notes"] or ["No problems found in the file."]})),
              ("robots.txt", pd.DataFrame({"Line": result["robots_text"].splitlines() or [""]}))]
    if draft is not None:
        sheets.append(("Draft vs live", draft))
    return checks.simple_excel("Robots.txt and AI crawlers", info, numbers, sheets)
